"""Planilhas dos setores (ANALISTAS, CONSULTORES TÉCNICOS e REGULADORES), abastecidas pelo sinistro.

Cada aba de trabalho é a aba SIGA do sinistro (colunas A a AO) filtrada por uma regra, mais uma ou
duas colunas próprias do setor:

    ANALISTAS.xlsx
        ANALISTAS            POR TAREFA = ANALISTAS, exceto VIDROS        + PARTICIPAÇÃO, OBSERVAÇÃO*
        VIDROS               Evento = VIDROS e SETOR diferente de FINANCEIRO + PARTICIPAÇÃO, OBSERVAÇÃO*
        OFICINAS             cópia da tabela Oficinas dos consultores
    CONSULTORES TECNICOS.xlsx
        CONSULTORES TEC.     POR TAREFA = CONSULTORES, exceto VIDROS      + CONSULTOR TECNICO, OBSERVAÇÃO*
        Oficinas             tabela digitada; o script preenche CNPJ LIMPO, QUANTIDADE e VAGAS DISPONIVEIS
        Oficinas por Consultor   oficinas com veículos, por consultor
    REGULADORES.xlsx
        REGULADORES          SETOR = REGULADORES, exceto VIDROS           + REGULARIZADO POR*
        VISTORIA INICIAL     tudo, exceto VIDROS                          + DATA DA REGULAGEM, MES REFERENCIA
        NEGADOS EM ABERTO    STATUS SGA = NEGADO e STATUS SIGA diferente de NEGADO
        NEGADOS FINALIZADOS  eventos do SIGA liberados no ano, com a situação do SGA
        HISTÓRICO REGULARIZAÇÕES   registro de cada REGULARIZADO POR preenchido

(*) coluna digitada pela equipe. O que for digitado é guardado em dados/manuais.json e recolocado na
linha do mesmo processo a cada atualização, mesmo que a linha mude de posição ou saia e volte.

Os processos das abas ANALISTAS e VIDROS também vão para a página ACOMPANHAMENTO ANALISTAS.html,
onde os analistas preenchem situação, próxima ação e prazo (ver acompanhamento.py).
"""
from __future__ import annotations

import collections
import copy
import datetime
import json
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from . import acompanhamento as A
from . import integracao as I
from . import planilha as P
from .util import data_br, texto, trim

N_SIGA = 41  # colunas A..AO da aba SIGA, repetidas em todas as abas de setor
AL_CNPJ, AP_PARTICIPACAO = I.AL, I.AP
DATAS_SIGA = (4, 12, 13, 17, 18, 19, I.Y, I.AA)  # E, M, N, R, S, T, Y, AA
MESES_SIGA = (I.A_COMPETENCIA, I.AI)             # Competência e MÊS REF SIGA
F_DIAS = '=IF(OR({c}{r}="",{c}{r}="00/00/0000"),"",IFERROR(NETWORKDAYS({c}{r},TODAY()),""))'
FMT_DATA, FMT_MES, FMT_DATA_HORA = "DD/MM/YYYY", "MM/YYYY", "DD/MM/YYYY HH:MM"

CAB_OFICINAS = ["OFICINAS", "TIPO", "CNPJ", "CNPJ LIMPO", "CONSULTOR", "ENDEREÇO", "CIDADE", "ESTADO", "SITUAÇÃO",
                "QUANTIDADE", "VAGAS", "VAGAS DISPONIVEIS"]
OF_NOME, OF_CNPJ, OF_LIMPO, OF_CONSULTOR, OF_UF, OF_QTD, OF_VAGAS, OF_DISPONIVEIS = 0, 2, 3, 4, 7, 9, 10, 11
CAB_OFICINAS_CONSULTOR = ["CONSULTOR", "OFICINA", "ESTADO", "QUANTIDADE"]
CAB_FINALIZADOS = ["Numero", "Origem", "STATUS SIGA", "Nome", "Cobertura", "Placa", "Abertura", "Liberação", "Duração",
                   "Oficina", "PROTOCOLO SGA", "STATUS SGA", "CLASSIFICAÇÃO NEGATIVA", "ESTADO", "COOPERATIVAS", "DATA EVENTO"]
CAB_BASE_VISTORIA = ["NUMERO", "REGULADOR ANALISTA", "DATA DA ANALISE"]
CAB_HISTORICO = ["DATA DA REGULARIZAÇÃO", "PROTOCOLO SIGA", "TAREFA", "REGULARIZADO POR"]
CAB_REG_ANALISTA = ["Nome", "Placa", "REGULADOR ANALISTA", "REGULADOR ORÇAMENTISTA", "DATA DA ANALISE", "MES REFERENCIA"]
CAB_VISTORIA_PRESENCIAL = ["SEM REGRESSO", "NOME", "PLACA", "ANALISTA", "REGULADOR", "DATA AGENDADO", "DATA REALIZADA",
                           "MÊS REFERENCIA", "OBS"]

ABA_ANALISTAS, ABA_VIDROS, ABA_OFICINAS_ESPELHO, ABA_VISTORIA_PRESENCIAL = "ANALISTAS", "VIDROS", "OFICINAS", "VISTORIA PRESENCIAL"
ABA_CONSULTORES, ABA_OFICINAS, ABA_OFICINAS_CONSULTOR = "CONSULTORES TEC.", "Oficinas", "Oficinas por Consultor"
ABA_REGULADORES, ABA_VISTORIA_INICIAL, ABA_NEGADOS_ABERTO, ABA_NEGADOS_FINALIZADOS = (
    "REGULADORES", "VISTORIA INICIAL", "NEGADOS EM ABERTO", "NEGADOS FINALIZADOS")
ABA_BASE_VISTORIA, ABA_HISTORICO, ABA_REG_ANALISTA = "BASE VISTORIA INICIAL", "HISTÓRICO REGULARIZAÇÕES", "REGULADORES ANALISTA"


@dataclass
class Linha:
    """Uma linha da aba SIGA do sinistro, já com as colunas calculadas."""
    v: list
    setor: str
    por_tarefa: str

    @property
    def numero(self) -> str:
        return trim(self.v[I.D_PROTOCOLO_SIGA])

    @property
    def tarefa(self) -> str:
        return trim(self.v[I.L_TAREFA])

    @property
    def vidros(self) -> bool:
        return trim(self.v[I.H_EVENTO]).upper() == "VIDROS"


def _mapa_tarefas(tarefas: list[list] | None) -> dict[str, tuple[str, str]]:
    """Aba TAREFAS: tarefa (maiúsculas) -> (SETOR, POR TAREFA)."""
    mapa: dict[str, tuple[str, str]] = {}
    for linha in (tarefas or [])[1:]:
        linha = list(linha) + [None] * 3
        if trim(linha[0]):
            mapa.setdefault(trim(linha[0]).upper(), (trim(linha[1]), trim(linha[2])))
    return mapa


class SetoresVistos:
    """Último setor em que cada processo apareceu na aba SIGA (dados/setores_vistos.json).

    A aba SIGA só tem os processos com tarefa aberta: quando o evento é finalizado, ele sai dela e o
    setor se perde. Anotar a cada execução o setor de cada processo que está na aba permite saber,
    depois, por qual setor ele passou por último (usado no filtro da conferência de finalizados).
    """

    def __init__(self, pasta_dados: Path):
        self.arquivo = pasta_dados / "setores_vistos.json"
        self.dados: dict[str, dict] = {}   # PROTOCOLO SIGA -> {"setores": [...], "tarefas": [...], "visto_em": aaaa-mm-dd}

    def carregar(self) -> None:
        if self.arquivo.exists():
            self.dados = json.loads(self.arquivo.read_text(encoding="utf-8"))

    def salvar(self) -> None:
        self.arquivo.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.arquivo.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.dados, ensure_ascii=False), encoding="utf-8")
        tmp.replace(self.arquivo)

    def registrar(self, siga: list[tuple[int, list]], tarefas: list[list] | None, quando: str) -> int:
        """Anota o setor de cada processo da aba SIGA (`quando` = aaaa-mm-dd). Devolve quantos processos havia."""
        mapa = _mapa_tarefas(tarefas)
        atuais: dict[str, tuple[set, set]] = collections.defaultdict(lambda: (set(), set()))
        for _, linha in siga:
            numero, tarefa = trim(linha[I.D_PROTOCOLO_SIGA]), trim(linha[I.L_TAREFA])
            if numero and tarefa:
                atuais[numero][0].add(mapa.get(tarefa.upper(), ("SETOR DESCONHECIDO",))[0] or "SETOR DESCONHECIDO")
                atuais[numero][1].add(tarefa)
        for numero, (setores_, tarefas_) in atuais.items():
            if self.dados.get(numero, {}).get("visto_em", "") <= quando:   # uma foto mais antiga não apaga a mais nova
                self.dados[numero] = {"setores": sorted(setores_), "tarefas": sorted(tarefas_), "visto_em": quando}
        return len(atuais)

    def setores(self, numero: str) -> list[str]:
        return self.dados.get(trim(numero), {}).get("setores", [])


def montar_linhas(siga: list[tuple[int, list]], resultado: I.Resultado, tarefas: list[list] | None) -> list[Linha]:
    """Junta o que está digitado na aba SIGA com as colunas calculadas e classifica cada tarefa pelo setor."""
    mapa = _mapa_tarefas(tarefas)
    saida = []
    for numero, atual in siga:
        v = [P._valor_modelo(x) for x in atual]
        for coluna, valor in resultado.siga[numero].items():
            v[coluna] = valor
        for coluna in DATAS_SIGA:  # datas de verdade, para ordenar e filtrar; "00/00/0000" continua texto
            v[coluna] = data_br(v[coluna]) or v[coluna]
        for coluna in MESES_SIGA:  # "09/2026" vira o dia 1º do mês, exibido como 09/2026, para ordenar certo
            v[coluna] = data_br("01/" + texto(v[coluna])) or v[coluna]
        tarefa = trim(v[I.L_TAREFA])
        setor, por_tarefa = mapa.get(tarefa.upper(), ("SETOR DESCONHECIDO",) * 2) if tarefa else ("SEM TAREFA",) * 2
        v[I.AJ], v[I.AQ] = setor, por_tarefa
        saida.append(Linha(v, setor, por_tarefa))
    return saida


def dias_uteis(inicio: datetime.date, fim: datetime.date) -> int:
    """Como o DIATRABALHOTOTAL (NETWORKDAYS) sem feriados: dias de segunda a sexta entre as datas, contando as duas."""
    if inicio > fim:
        return -dias_uteis(fim, inicio)
    semanas, resto = divmod((fim - inicio).days + 1, 7)
    return semanas * 5 + sum(1 for k in range(resto) if (inicio.weekday() + k) % 7 < 5)


def _dias_ate_hoje(v, hoje: datetime.date):
    d = v if isinstance(v, datetime.date) else data_br(v)
    return dias_uteis(d, hoje) if d else ""


def _linhas_aba(linhas: list[Linha], filtro, extras, valores: bool = False) -> list[list]:
    """Linhas de uma aba de setor. Os dias úteis (Z e AB) vão como fórmula (Excel, contam sozinhos a cada dia)
    ou, com `valores`, já calculados para hoje (Google, onde o sistema não grava fórmulas)."""
    saida = []
    hoje = datetime.date.today()
    for linha in linhas:
        if filtro(linha):
            r = len(saida) + 2
            v = list(linha.v[:N_SIGA])
            if valores:
                v[I.Z], v[I.AB] = _dias_ate_hoje(v[I.Y], hoje), _dias_ate_hoje(v[I.AA], hoje)
            else:
                v[I.Z] = P.Formula(F_DIAS.format(c="Y", r=r))
                v[I.AB] = P.Formula(F_DIAS.format(c="AA", r=r))
            saida.append(v + extras(linha))
    return saida


def _digitos(v) -> str:
    return re.sub(r"\D", "", texto(v))


# ----------------------------------------------------------------------------- o que a equipe digita

class Manuais:
    """Valores digitados nas colunas manuais, guardados por processo em dados/manuais.json."""

    def __init__(self, pasta_dados: Path):
        self.arquivo = pasta_dados / "manuais.json"
        self.arquivo_gravados = pasta_dados / "manuais_gravados.json"
        self.dados: dict[str, dict[str, str]] = {}
        # O que o sistema gravou por último na planilha, por grupo: comparando com isso se sabe o que a
        # equipe editou nela, mesmo quando o valor guardado mudou por outro caminho (a página de acompanhamento).
        self.gravados: dict[str, dict[str, str]] = {}

    def carregar(self) -> None:
        if self.arquivo.exists():
            self.dados = json.loads(self.arquivo.read_text(encoding="utf-8"))
        if self.arquivo_gravados.exists():
            self.gravados = json.loads(self.arquivo_gravados.read_text(encoding="utf-8"))
        else:   # até aqui o que estava guardado era sempre o que tinha sido gravado na planilha
            self.gravados = copy.deepcopy(self.dados)

    def salvar(self) -> None:
        self.arquivo.parent.mkdir(parents=True, exist_ok=True)
        for arquivo, conteudo in ((self.arquivo, self.dados), (self.arquivo_gravados, self.gravados)):
            tmp = arquivo.with_suffix(".tmp")
            tmp.write_text(json.dumps(conteudo, ensure_ascii=False, indent=1), encoding="utf-8")
            tmp.replace(arquivo)

    def obter(self, grupo: str, chave: str) -> str:
        return self.dados.get(grupo, {}).get(chave, "")

    def referencia(self, grupo: str) -> dict[str, str]:
        return self.gravados.get(grupo, {})

    def gravou(self, grupo: str, chaves) -> None:
        """Anota o que acabou de ser gravado na planilha para estas chaves."""
        self.gravados[grupo] = {c: self.obter(grupo, c) for c in chaves if c and self.obter(grupo, c)}

    def colher(self, grupo: str, pares) -> list[tuple[str, str]]:
        """Guarda o que está na planilha. Devolve os (chave, valor) novos ou alterados."""
        guardado = self.dados.setdefault(grupo, {})
        mudou = []
        for chave, valor in pares:
            if not chave:
                continue
            if valor:
                if guardado.get(chave) != valor:
                    mudou.append((chave, valor))
                    guardado[chave] = valor
            else:
                guardado.pop(chave, None)  # a equipe apagou a anotação
        return mudou

    def colher_edicoes(self, grupo: str, pares, referencia: dict[str, str]) -> list[str]:
        """Guarda só o que foi editado desde a última gravação (`referencia` = o que o sistema gravou).

        Serve para a mesma anotação aparecer em mais de uma aba: vale a aba onde ela foi alterada.
        Devolve as chaves alteradas.
        """
        guardado = self.dados.setdefault(grupo, {})
        mudou = []
        for chave, valor in pares:
            if not chave or valor == referencia.get(chave, ""):
                continue
            if valor:
                guardado[chave] = valor
            else:
                guardado.pop(chave, None)
            mudou.append(chave)
        return mudou


def _pares(tabela: list[list] | None, coluna: int, com_tarefa: bool = False):
    """(chave, valor digitado) de cada linha de uma aba de setor. Chave = PROTOCOLO SIGA [|Tarefa]."""
    for linha in (tabela or [])[1:]:
        linha = list(linha) + [None] * (coluna + 1)
        chave = trim(linha[I.D_PROTOCOLO_SIGA])
        if chave and com_tarefa:
            chave = f"{chave}|{trim(linha[I.L_TAREFA])}"
        yield chave, texto(linha[coluna]).strip()


# ----------------------------------------------------------------------------- leitura das planilhas originais

def _ler_base(caminho: Path | None, pedidos: dict[str, object], log) -> dict[str, list[list] | None]:
    """Lê abas da planilha original do Google. Cada pedido é um nome exato ou uma expressão regular;
    quando há várias abas com o padrão (as cópias diárias), vale a que está visível."""
    vazio = {apelido: None for apelido in pedidos}
    if not caminho:
        return vazio
    if not caminho.exists():
        log.warning("Planilha original %s não encontrada; a planilha nova começa sem o histórico dela.", caminho)
        return vazio
    log.info("Primeira carga: lendo %s (pode levar alguns minutos)...", caminho.name)
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        visiveis = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        saida = {}
        for apelido, alvo in pedidos.items():
            if hasattr(alvo, "fullmatch"):
                teste = lambda n, alvo=alvo: alvo.fullmatch(n.strip()) is not None
            else:
                teste = lambda n, alvo=alvo: n.strip() == alvo
            nome = next((n for n in visiveis if teste(n)), None) or next((n for n in wb.sheetnames if teste(n)), None)
            if nome is None:
                saida[apelido] = None
                continue
            ws = wb[nome]
            ws.reset_dimensions()
            saida[apelido] = [list(l) for l in ws.iter_rows(values_only=True)]
        return saida
    finally:
        wb.close()


def _dados(tabela: list[list] | None, n_colunas: int) -> list[list]:
    """Linhas preenchidas (sem o cabeçalho), cortadas em n_colunas."""
    saida = []
    for linha in (tabela or [])[1:]:
        linha = (list(linha) + [None] * n_colunas)[:n_colunas]
        if any(texto(x).strip() for x in linha):
            saida.append(linha)
    return saida


def _titulos(tabela: list[list] | None, padrao: list[str], ate: int | None = None) -> list[str]:
    """Títulos da planilha original (as equipes já conhecem), completados com os nomes padrão."""
    nomes = list(padrao)
    if tabela:
        for c, v in enumerate(tabela[0][:ate if ate is not None else len(padrao)]):
            if c < len(nomes) and trim(v):
                nomes[c] = trim(v)
    return nomes


# ----------------------------------------------------------------------------- criação das planilhas

def _aba(wb: Workbook, nome: str, cabecalho: list[str], manuais: tuple = (), formatos: dict | None = None,
         linhas: list[list] | None = None):
    """Aba nova. Sem `linhas`, é uma aba gerada (só cabeçalho + linha-molde de estilo). `manuais` são as
    colunas digitadas pela equipe (título azul); as demais ficam com título verde."""
    ws = wb.create_sheet(nome)
    P.montar_cabecalho(ws, cabecalho, calculadas=[c for c in range(len(cabecalho)) if c not in manuais])
    formatos = formatos or {}
    ultima = 1
    for linha in linhas or []:
        ultima += 1
        for c, valor in enumerate(linha):
            valor = P._valor_modelo(valor)
            if valor is None or valor == "":
                continue
            cel = ws.cell(row=ultima, column=c + 1, value=valor)
            if c in formatos:
                cel.number_format = formatos[c]
            elif isinstance(valor, (datetime.date, datetime.datetime)):
                cel.number_format = FMT_DATA
    if ultima == 1:
        for c, formato in formatos.items():
            ws.cell(row=2, column=c + 1).number_format = formato
        ultima = 2
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecalho))}{ultima}"
    return ws


def _formatos_setor(extras: dict | None = None) -> dict:
    formatos = {c: FMT_DATA for c in DATAS_SIGA}
    formatos.update({c: FMT_MES for c in MESES_SIGA})
    formatos.update(extras or {})
    return formatos


def _salvar_nova(wb: Workbook, destino: Path) -> None:
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])  # a aba vazia que o openpyxl cria por padrão
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


# ----------------------------------------------------------------------------- contexto de uma execução

@dataclass
class Contexto:
    cfg: object
    log: object
    linhas: list[Linha]
    manuais: Manuais
    resultado: I.Resultado
    sga: list[list[str]]
    integ: list[list]
    historico: A.Historico | None = None

    def abrir(self, destino: Path, criar) -> tuple[P.EditorXlsx, bool]:
        """Garante que a planilha existe (criando na primeira vez) e devolve (editor, é nova?)."""
        nova = not destino.exists()
        if nova:
            criar()
            self.log.info("Planilha criada: %s", destino)
        elif P.planilha_bloqueada(destino):
            raise P.PlanilhaAberta(f"A planilha {destino} está aberta em outro programa. Feche-a e rode de novo.")
        return P.EditorXlsx(destino), nova

    def salvar(self, editor: P.EditorXlsx, nova: bool, resumo: str) -> None:
        backup = P.salvar_editor(editor, self.cfg.pasta_backups, self.cfg.backups_manter, fazer_backup=not nova)
        self.log.info("%s: %s%s", editor.caminho.name, resumo, f" (backup: {backup})" if backup else "")


# ----------------------------------------------------------------------------- CONSULTORES TÉCNICOS

def eh_consultores(l: Linha) -> bool:
    return l.por_tarefa == "CONSULTORES" and not l.vidros


def tabela_oficinas(tabela: list[list] | None) -> list[tuple[int, list]]:
    """(nº da linha, valores) das oficinas preenchidas da aba Oficinas."""
    return [(r, (list(l) + [None] * len(CAB_OFICINAS))[:len(CAB_OFICINAS)])
            for r, l in enumerate(tabela or [], 1)
            if r > 1 and (trim(l[OF_NOME] if l else None) or _digitos(l[OF_CNPJ] if len(l) > OF_CNPJ else None))]


def consultores_por_cnpj(oficinas: list[tuple[int, list]]) -> dict[str, str]:
    mapa: dict[str, str] = {}
    for _, of in oficinas:
        if _digitos(of[OF_CNPJ]):
            mapa.setdefault(_digitos(of[OF_CNPJ]), trim(of[OF_CONSULTOR]))
    return mapa


def consultor_da_linha(linha: Linha, por_cnpj: dict[str, str]) -> str:
    cnpj = _digitos(linha.v[AL_CNPJ])
    return por_cnpj.get(cnpj, "NÃO ENCONTRADO") if cnpj else "NÃO ENCONTRADO"


def _consultores(ctx: Contexto) -> list[list] | None:
    """Atualiza a planilha dos consultores técnicos. Devolve a tabela de oficinas já calculada."""
    destino = ctx.cfg.planilha_consultores
    grupo = f"{ABA_CONSULTORES}/OBSERVAÇÃO"
    cab_setor = P.CABECALHO_SIGA[:N_SIGA] + ["CONSULTOR TECNICO", "OBSERVAÇÃO"]

    def criar():
        base = _ler_base(ctx.cfg.base_consultores, {"setor": ABA_CONSULTORES, "oficinas": ABA_OFICINAS}, ctx.log)
        ctx.manuais.colher(grupo, _pares(base["setor"], N_SIGA + 1))
        wb = Workbook()
        _aba(wb, ABA_CONSULTORES, _titulos(base["setor"], cab_setor, N_SIGA), manuais=(N_SIGA + 1,), formatos=_formatos_setor())
        _aba(wb, ABA_OFICINAS, _titulos(base["oficinas"], CAB_OFICINAS),
             manuais=tuple(c for c in range(len(CAB_OFICINAS)) if c not in (OF_LIMPO, OF_QTD, OF_DISPONIVEIS)),
             linhas=_dados(base["oficinas"], len(CAB_OFICINAS)))
        _aba(wb, ABA_OFICINAS_CONSULTOR, CAB_OFICINAS_CONSULTOR)
        _salvar_nova(wb, destino)

    editor, nova = ctx.abrir(destino, criar)
    atual = P.ler_abas(destino, [ABA_CONSULTORES, ABA_OFICINAS])
    if not nova:
        ctx.manuais.colher(grupo, _pares(atual[ABA_CONSULTORES], N_SIGA + 1))

    oficinas = tabela_oficinas(atual[ABA_OFICINAS])
    por_cnpj = consultores_por_cnpj(oficinas)
    linhas = _linhas_aba(ctx.linhas, eh_consultores, lambda l: [consultor_da_linha(l, por_cnpj), ctx.manuais.obter(grupo, l.numero)])
    editor.substituir_aba(ABA_CONSULTORES, linhas, len(cab_setor))

    veiculos = collections.Counter(_digitos(l[AL_CNPJ]) for l in linhas)
    alteracoes, tabela, resumo = {}, [], []
    for r, of in oficinas:
        limpo = _digitos(of[OF_CNPJ])
        quantidade = veiculos.get(limpo, 0) if limpo else 0
        alteracoes[r] = {OF_LIMPO: limpo, OF_QTD: quantidade, OF_DISPONIVEIS: P.Formula(f'=IFERROR(K{r}-J{r},"")')}
        vagas = texto(of[OF_VAGAS]).strip()
        of = list(of)
        of[OF_LIMPO], of[OF_QTD] = limpo, quantidade
        of[OF_DISPONIVEIS] = int(vagas) - quantidade if re.fullmatch(r"\d+", vagas) else ""
        tabela.append(of)
        if quantidade > 0:
            resumo.append([trim(of[OF_CONSULTOR]), trim(of[OF_NOME]), trim(of[OF_UF]), quantidade])
    editor.atualizar_aba(ABA_OFICINAS, alteracoes)
    resumo.sort(key=lambda x: (x[0], -x[3], x[1]))
    editor.substituir_aba(ABA_OFICINAS_CONSULTOR, resumo, len(CAB_OFICINAS_CONSULTOR))

    sem_consultor = sum(1 for l in linhas if l[N_SIGA] == "NÃO ENCONTRADO")
    ctx.salvar(editor, nova, f"{ABA_CONSULTORES} {len(linhas)} linhas ({sem_consultor} sem consultor técnico); {len(tabela)} oficinas")
    return tabela


# ----------------------------------------------------------------------------- ANALISTAS

def _eh_analistas(l: Linha) -> bool:
    return l.por_tarefa == "ANALISTAS" and not l.vidros


def _eh_vidros(l: Linha) -> bool:
    return l.vidros and l.setor != "FINANCEIRO"


def processos_analistas(linhas: list[Linha]) -> list[tuple[str, list[Linha]]]:
    """Os processos das abas ANALISTAS e VIDROS (também usados na página de acompanhamento)."""
    return [(ABA_ANALISTAS, [l for l in linhas if _eh_analistas(l)]), (ABA_VIDROS, [l for l in linhas if _eh_vidros(l)])]


def _analistas(ctx: Contexto, oficinas: list[list] | None) -> None:
    destino = ctx.cfg.planilha_analistas
    grupos = {ABA_ANALISTAS: f"{ABA_ANALISTAS}/OBSERVAÇÃO", ABA_VIDROS: f"{ABA_VIDROS}/OBSERVAÇÃO"}
    cab_setor = P.CABECALHO_SIGA[:N_SIGA] + ["PARTICIPAÇÃO", "OBSERVAÇÃO"]

    def criar():
        base = _ler_base(ctx.cfg.base_analistas, {
            ABA_ANALISTAS: re.compile(r"\d{4}\s*ANALISTAS", re.I), ABA_VIDROS: re.compile(r"VIDROS\s*\d{4}", re.I),
            "presencial": ABA_VISTORIA_PRESENCIAL}, ctx.log)
        wb = Workbook()
        for aba, grupo in grupos.items():
            ctx.manuais.colher(grupo, _pares(base[aba], N_SIGA + 1))
            _aba(wb, aba, _titulos(base[aba], cab_setor, N_SIGA), manuais=(N_SIGA + 1,), formatos=_formatos_setor())
        _aba(wb, ABA_OFICINAS_ESPELHO, CAB_OFICINAS)
        presencial = _dados(base["presencial"], len(CAB_VISTORIA_PRESENCIAL))
        for r, linha in enumerate(presencial, 2):
            linha[7] = f"=G{r}" if linha[6] not in (None, "") else None
        _aba(wb, ABA_VISTORIA_PRESENCIAL, _titulos(base["presencial"], CAB_VISTORIA_PRESENCIAL),
             manuais=tuple(c for c in range(len(CAB_VISTORIA_PRESENCIAL)) if c != 7), formatos={7: FMT_MES}, linhas=presencial)
        _salvar_nova(wb, destino)

    editor, nova = ctx.abrir(destino, criar)
    if not nova:
        # Vale só o que foi editado na planilha desde a última gravação: o mesmo campo pode ter sido
        # alterado pela página de acompanhamento, e a planilha ainda mostra o valor antigo.
        atual = P.ler_abas(destino, list(grupos))
        for aba, grupo in grupos.items():
            referencia = ctx.manuais.referencia(grupo)
            mudou = ctx.manuais.colher_edicoes(grupo, _pares(atual[aba], N_SIGA + 1), referencia)
            if mudou and ctx.historico is not None:
                A.anotar_planilha(ctx.manuais, ctx.historico, aba,
                                  [(c, referencia.get(c, ""), ctx.manuais.obter(grupo, c)) for c in mudou])

    def extras(grupo):
        return lambda l: [l.v[AP_PARTICIPACAO], ctx.manuais.obter(grupo, l.numero)]

    analistas = _linhas_aba(ctx.linhas, _eh_analistas, extras(grupos[ABA_ANALISTAS]))
    vidros = _linhas_aba(ctx.linhas, _eh_vidros, extras(grupos[ABA_VIDROS]))
    editor.substituir_aba(ABA_ANALISTAS, analistas, len(cab_setor))
    editor.substituir_aba(ABA_VIDROS, vidros, len(cab_setor))
    if oficinas is not None:
        editor.substituir_aba(ABA_OFICINAS_ESPELHO, oficinas, len(CAB_OFICINAS))
    ctx.salvar(editor, nova, f"{ABA_ANALISTAS} {len(analistas)} linhas; {ABA_VIDROS} {len(vidros)} linhas")
    for aba, lista in processos_analistas(ctx.linhas):
        ctx.manuais.gravou(grupos[aba], (l.numero for l in lista))


# ----------------------------------------------------------------------------- REGULADORES

def _negados_finalizados(ctx: Contexto) -> list[list]:
    """Eventos do SIGA com Liberação no ano corrente, com a situação do evento no SGA."""
    ano = datetime.date.today().year
    linhas = I.finalizados(ctx.sga, ctx.integ, ctx.resultado, datetime.date(ano, 1, 1), datetime.date(ano, 12, 31))
    for linha in linhas:  # Abertura e Liberação como datas de verdade na planilha
        linha[6], linha[7] = data_br(linha[6]) or linha[6], data_br(linha[7]) or linha[7]
    return linhas


def eh_reguladores(l: Linha) -> bool:
    return l.setor == "REGULADORES" and not l.vidros


def eh_vistoria(l: Linha) -> bool:
    return not l.vidros


def eh_negado_aberto(l: Linha) -> bool:
    return trim(l.v[I.AE]).upper() == "NEGADO" and trim(l.v[I.AG]).upper() != "NEGADO"


def regulagens(tabela: list[list] | None) -> dict[str, object]:
    """BASE VISTORIA INICIAL (NUMERO, REGULADOR ANALISTA, DATA DA ANALISE) -> {número: data da análise}."""
    regulagem: dict[str, object] = {}
    for linha in _dados(tabela, 3):
        if trim(linha[0]):
            regulagem.setdefault(trim(linha[0]), A._como_data(linha[2]) or texto(linha[2]))
    return regulagem


def extras_vistoria(linha: Linha, regulagem: dict[str, object]) -> list:
    data = regulagem.get(linha.numero, "")
    return [data, data.replace(day=1) if isinstance(data, datetime.date) else ""]


def _reguladores(ctx: Contexto) -> None:
    destino = ctx.cfg.planilha_reguladores
    grupo = f"{ABA_REGULADORES}/REGULARIZADO POR"
    cab_reguladores = P.CABECALHO_SIGA[:N_SIGA] + ["REGULARIZADO POR"]
    cab_vistoria = P.CABECALHO_SIGA[:N_SIGA] + ["DATA DA REGULAGEM", "MES REFERENCIA"]
    cab_negados = P.CABECALHO_SIGA[:N_SIGA]

    def criar():
        base = _ler_base(ctx.cfg.base_reguladores, {
            ABA_REGULADORES: ABA_REGULADORES, ABA_VISTORIA_INICIAL: ABA_VISTORIA_INICIAL, ABA_NEGADOS_ABERTO: ABA_NEGADOS_ABERTO,
            "controle": "_CONTROLE_AP", ABA_HISTORICO: ABA_HISTORICO, ABA_REG_ANALISTA: ABA_REG_ANALISTA}, ctx.log)
        sinistro = _ler_base(ctx.cfg.planilha_base, {ABA_BASE_VISTORIA: ABA_BASE_VISTORIA}, ctx.log)
        # O que já estava regularizado entra sem gerar novo registro no histórico.
        ctx.manuais.colher(grupo, ((f"{trim(l[0])}|{trim(l[1])}", trim(l[2])) for l in _dados(base["controle"], 3)))
        ctx.manuais.colher(grupo, _pares(base[ABA_REGULADORES], N_SIGA, com_tarefa=True))
        wb = Workbook()
        _aba(wb, ABA_REGULADORES, _titulos(base[ABA_REGULADORES], cab_reguladores, N_SIGA), manuais=(N_SIGA,), formatos=_formatos_setor())
        _aba(wb, ABA_VISTORIA_INICIAL, _titulos(base[ABA_VISTORIA_INICIAL], cab_vistoria, N_SIGA),
             formatos=_formatos_setor({N_SIGA: FMT_DATA, N_SIGA + 1: FMT_MES}))
        _aba(wb, ABA_NEGADOS_ABERTO, _titulos(base[ABA_NEGADOS_ABERTO], cab_negados, N_SIGA), formatos=_formatos_setor())
        _aba(wb, ABA_NEGADOS_FINALIZADOS, CAB_FINALIZADOS, formatos={6: FMT_DATA, 7: FMT_DATA})
        _aba(wb, ABA_BASE_VISTORIA, CAB_BASE_VISTORIA, manuais=(0, 1, 2), formatos={2: FMT_DATA},
             linhas=_dados(sinistro[ABA_BASE_VISTORIA], len(CAB_BASE_VISTORIA)))
        _aba(wb, ABA_HISTORICO, CAB_HISTORICO, formatos={0: FMT_DATA_HORA}, linhas=_dados(base[ABA_HISTORICO], len(CAB_HISTORICO)))
        analista = _dados(base[ABA_REG_ANALISTA], len(CAB_REG_ANALISTA))
        for r, linha in enumerate(analista, 2):
            linha[5] = f'=IFERROR(IF(E{r}="","",DATE(YEAR(E{r}),MONTH(E{r}),1)),"")'
        _aba(wb, ABA_REG_ANALISTA, CAB_REG_ANALISTA, manuais=(0, 1, 2, 3, 4), formatos={4: FMT_DATA, 5: FMT_MES}, linhas=analista)
        _salvar_nova(wb, destino)

    editor, nova = ctx.abrir(destino, criar)
    atual = P.ler_abas(destino, [ABA_REGULADORES, ABA_BASE_VISTORIA, ABA_HISTORICO])
    regularizados = [] if nova else ctx.manuais.colher(grupo, _pares(atual[ABA_REGULADORES], N_SIGA, com_tarefa=True))

    regulagem = regulagens(atual[ABA_BASE_VISTORIA])
    reguladores = _linhas_aba(ctx.linhas, eh_reguladores, lambda l: [ctx.manuais.obter(grupo, f"{l.numero}|{l.tarefa}")])
    vistoria = _linhas_aba(ctx.linhas, eh_vistoria, lambda l: extras_vistoria(l, regulagem))
    abertos = _linhas_aba(ctx.linhas, eh_negado_aberto, lambda l: [])
    finalizados = _negados_finalizados(ctx)
    editor.substituir_aba(ABA_REGULADORES, reguladores, len(cab_reguladores))
    editor.substituir_aba(ABA_VISTORIA_INICIAL, vistoria, len(cab_vistoria))
    editor.substituir_aba(ABA_NEGADOS_ABERTO, abertos, len(cab_negados))
    editor.substituir_aba(ABA_NEGADOS_FINALIZADOS, finalizados, len(CAB_FINALIZADOS))

    if regularizados:  # cada regularização nova vira uma linha no histórico, com a hora desta execução
        ultima = max((r for r, l in enumerate(atual[ABA_HISTORICO] or [], 1) if any(texto(x).strip() for x in l[:4])), default=1)
        agora = datetime.datetime.now().replace(microsecond=0)
        registro = {}
        for k, (chave, quem) in enumerate(regularizados, 1):
            numero, _, tarefa = chave.partition("|")
            registro[ultima + k] = {0: agora, 1: numero, 2: tarefa, 3: quem}
        editor.atualizar_aba(ABA_HISTORICO, registro)

    ctx.salvar(editor, nova, f"{ABA_REGULADORES} {len(reguladores)}; {ABA_VISTORIA_INICIAL} {len(vistoria)} "
               f"({sum(1 for l in vistoria if l[N_SIGA])} com regulagem); {ABA_NEGADOS_ABERTO} {len(abertos)}; "
               f"{ABA_NEGADOS_FINALIZADOS} {len(finalizados)}; {len(regularizados)} regularizações novas no histórico")


# ----------------------------------------------------------------------------- entrada

def atualizar(cfg, log, siga: list[tuple[int, list]], resultado: I.Resultado, tarefas: list[list] | None,
              sga: list[list[str]], integ: list[list], edicoes: list[Path] | None = None, google=None,
              base_vistoria: list[list] | None = None) -> tuple[int, list[Path]]:
    """Incorpora o que foi salvo pela página de acompanhamento (`edicoes`), gera/atualiza as planilhas de
    setor configuradas e refaz a página. Devolve (quantas etapas falharam, arquivos de `edicoes` incorporados).

    `google` (google_planilhas.ContaGoogle): a planilha de setor configurada no Google é gravada lá, no lugar do
    .xlsx; `base_vistoria` é a aba BASE VISTORIA INICIAL do sinistro no Google (datas da regulagem)."""
    from . import setores_google as SG
    no_google = {chave: bool(google and getattr(cfg, f"google_{chave}")) for chave in ("consultores", "analistas", "reguladores")}
    if not (cfg.planilha_consultores or cfg.planilha_analistas or cfg.planilha_reguladores or cfg.pagina_acompanhamento
            or any(no_google.values())):
        return 0, []
    manuais = Manuais(cfg.pasta_dados)
    manuais.carregar()
    historico = A.Historico(cfg.pasta_dados)
    historico.carregar()
    incorporados = A.aplicar(edicoes or [], manuais, historico, cfg.registro, log)
    ctx = Contexto(cfg, log, montar_linhas(siga, resultado, tarefas), manuais, resultado, sga, integ, historico)
    desconhecidas = collections.Counter(l.tarefa for l in ctx.linhas if l.setor == "SETOR DESCONHECIDO")
    if desconhecidas:
        log.warning("Tarefas sem setor na aba TAREFAS (não entram em nenhuma planilha de setor): %s",
                    ", ".join(f"{t} ({n})" for t, n in desconhecidas.most_common()))
    falhas = 0
    oficinas = None
    etapas = []   # (rótulo para o log, etapa)
    if no_google["consultores"]:
        etapas.append(("CONSULTORES (Google)", lambda: SG.consultores(ctx, google.planilha(cfg.google_consultores))))
    elif cfg.planilha_consultores:
        etapas.append((cfg.planilha_consultores.name, lambda: _consultores(ctx)))
    if no_google["analistas"]:
        etapas.append(("ANALISTAS (Google)", lambda: SG.analistas(ctx, google.planilha(cfg.google_analistas))))
    elif cfg.planilha_analistas:
        etapas.append((cfg.planilha_analistas.name, lambda: _analistas(ctx, oficinas)))
    if no_google["reguladores"]:
        etapas.append(("REGULADORES (Google)", lambda: SG.reguladores(ctx, google.planilha(cfg.google_reguladores), base_vistoria)))
    elif cfg.planilha_reguladores:
        etapas.append((cfg.planilha_reguladores.name, lambda: _reguladores(ctx)))
    for rotulo, etapa in etapas:
        try:
            retorno = etapa()
            if cfg.planilha_consultores and rotulo == cfg.planilha_consultores.name:
                oficinas = retorno
        except SG.GoogleErro as e:
            log.error("%s: %s", rotulo, e)
            falhas += 1
        except P.PlanilhaAberta as e:
            log.error("%s", e)
            falhas += 1
        except KeyError as e:
            log.error("%s: %s Para recriá-la, renomeie o arquivo atual e rode de novo.", rotulo, str(e).strip("'\""))
            falhas += 1
        except Exception:
            log.exception("Erro ao atualizar %s; ela não foi alterada.", rotulo)
            falhas += 1
    manuais.salvar()
    historico.salvar()

    if cfg.pagina_acompanhamento:
        try:
            n = A.gerar_pagina(cfg.pagina_acompanhamento, processos_analistas(ctx.linhas), manuais, historico)
            log.info("Página %s: %d processos; %d a preencher, %d atrasados, %d em dia, %d concluídos.",
                     cfg.pagina_acompanhamento.name, sum(n.values()), n["PREENCHER"], n["ATRASADO"], n["EM DIA"], n["CONCLUÍDO"])
        except OSError as e:
            log.error("Não consegui gravar a página %s (%s).", cfg.pagina_acompanhamento, e)
            falhas += 1
    return falhas, incorporados
