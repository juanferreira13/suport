"""Leitura e gravação do .xlsx do sinistro.

A planilha é atualizada "por dentro" do arquivo: só o conteúdo das abas SGA e INTEGRAÇÃO e as
colunas calculadas da aba SIGA são reescritos. Todo o resto (o que foi digitado na aba SIGA,
formatação, outras abas, tabelas dinâmicas, gráficos) permanece exatamente como estava.
"""
from __future__ import annotations

import collections
import datetime
import html
import re
import shutil
import warnings
import zipfile
from pathlib import Path
from xml.parsers import expat

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from . import integracao as I
from . import sga_base as S
from . import siga as G
from .util import chave_cabecalho, texto

ABA_SGA, ABA_INTEGRACAO, ABA_SIGA, ABA_TAREFAS = "SGA", "INTEGRAÇÃO", "SIGA", "TAREFAS"

CABECALHO_SGA = [c for c, _ in S.COLUNAS] + ["NOME", "PLACA", "COMPETENCIA", "STATUS SIGA", "chave"]
CABECALHO_INTEGRACAO = G.COLUNAS + ["PROTOCOLO", "PLACA 2", "COMPETENCIA", "chave"]
CABECALHO_SIGA = [
    "Competência", "1", "#", "PROTOCOLO SIGA", "Data abertura", "Dias em aberto", "Cobertura", "Evento", "Nome",
    "Veículo", "Placa SIGA", "Tarefa", "Data início", "Data término", "Dias em atraso", "Responsável", "Oficina",
    "Data entrada oficina", "Data início reparo", "Data liberação", "Total rateio", "Tipo Veículo", "Associação",
    "ANALISTA", "DATA COMUNICADO EVENTO", "DIAS UTEIS COMUNICADO", "DATA AUTORIZADO", "DIAS UTEIS AUTORIZADO",
    "PLACA SGA", "ESTADO", "STATUS SGA", "PROTOCOLO SGA", "STATUS SIGA", "CLASSIFICAÇÃO NEGATIVA", "MÊS REF SIGA",
    "SETOR", "OFICINA SGA", "CNPJ", "Regulador", "VOLUNTARIO", "COOPERATIVA", "PARTICIPAÇÃO", "POR TAREFA",
    "CIDADE", "DATA EVENTO",
]
CABECALHO_TAREFAS = ["TAREFAS", "SETOR", "POR TAREFA"]

_LETRAS = [get_column_letter(i) for i in range(1, 200)]
_EPOCA = datetime.date(1899, 12, 30)
_ILEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")
_LINHA = re.compile(r"<row\b([^>]*?)/>|<row\b([^>]*)>(.*?)</row>", re.S)
_CELULA = re.compile(r"<c\b([^>]*?)/>|<c\b([^>]*)>(.*?)</c>", re.S)
_REF = re.compile(r'\br="([A-Z]+)(\d+)"')
_ESTILO = re.compile(r'\bs="(\d+)"')
_CONTEUDO = re.compile(r"<f\b|<is\b|<v>[^<]")
_ATRIBUTO = re.compile(r'([\w:]+)="([^"]*)"')


class PlanilhaAberta(Exception):
    """A planilha está aberta no Excel (ou sem permissão de escrita)."""


class Formula:
    def __init__(self, texto_formula: str):
        self.texto = texto_formula.lstrip("=")


class SeVazia:
    """Valor gravado apenas quando a célula ainda não tem nada (não sobrescreve o que já existe)."""

    def __init__(self, valor):
        self.valor = valor


# ----------------------------------------------------------------------------- leitura

def _achar_aba(nomes: list[str], procurada: str) -> str | None:
    alvo = chave_cabecalho(procurada)
    return next((n for n in nomes if chave_cabecalho(n) == alvo), None)


def ler_abas(caminho: Path, nomes: list[str | None], formulas: bool = False) -> dict:
    """Lê abas inteiras como listas de linhas. Nome None = primeira aba. Aba inexistente = None."""
    with warnings.catch_warnings():  # listas suspensas gravadas pelo Excel: o aviso do openpyxl não afeta a leitura
        warnings.filterwarnings("ignore", message="Data Validation extension")
        wb = load_workbook(caminho, read_only=True, data_only=not formulas)
        try:
            saida = {}
            for nome in nomes:
                real = wb.sheetnames[0] if nome is None else _achar_aba(wb.sheetnames, nome)
                if real is None:
                    saida[nome] = None
                    continue
                ws = wb[real]
                ws.reset_dimensions()  # não confia no tamanho declarado no arquivo; lê tudo
                saida[nome] = [list(l) for l in ws.iter_rows(values_only=True)]
            return saida
        finally:
            wb.close()


def linhas_siga(tabela: list[list] | None) -> list[tuple[int, list]]:
    """(nº da linha na planilha, valores A..AS) das linhas da aba SIGA que têm PROTOCOLO SIGA."""
    saida = []
    for numero, linha in enumerate(tabela or [], 1):
        linha = (list(linha) + [None] * I.SIGA_COLUNAS)[:I.SIGA_COLUNAS]
        if numero > 1 and texto(linha[I.D_PROTOCOLO_SIGA]).strip():
            saida.append((numero, linha))
    return saida


# ----------------------------------------------------------------------------- criação do modelo

def _valor_modelo(v):
    if isinstance(v, datetime.datetime) and v.time() == datetime.time(0):
        return v.date()
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def montar_cabecalho(ws, nomes, calculadas=()) -> None:
    """Linha de títulos: azul para colunas digitadas, verde para as que o script preenche."""
    titulo = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F4E78")
    fundo_calc = PatternFill("solid", fgColor="548235")
    centro = Alignment(vertical="center", wrap_text=True)
    for c, nome in enumerate(nomes, 1):
        cel = ws.cell(row=1, column=c, value=nome)
        cel.font, cel.alignment = titulo, centro
        cel.fill = fundo_calc if c - 1 in calculadas else fundo
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(45, len(nome) + 6))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def criar_modelo(destino: Path, siga_valores: list[list] | None, tarefas: list[list] | None) -> None:
    """Cria a planilha do sinistro com as abas TAREFAS, SGA, INTEGRAÇÃO e SIGA.

    As abas SGA e INTEGRAÇÃO nascem só com o cabeçalho (são preenchidas na gravação). A aba SIGA
    recebe os valores da planilha de origem; a TAREFAS é copiada.
    """
    cabecalho = montar_cabecalho
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_TAREFAS
    cabecalho(ws, CABECALHO_TAREFAS)
    for linha in (tarefas or [])[1:]:
        if texto(linha[0] if linha else None).strip():
            ws.append([_valor_modelo(v) for v in linha[:3]])
    ws.column_dimensions["A"].width = 48

    # A linha 2 fica formatada (texto/data/número) e serve de molde de estilo para a gravação.
    ws = wb.create_sheet(ABA_SGA)
    cabecalho(ws, CABECALHO_SGA, calculadas=range(S.N_COLUNAS, len(CABECALHO_SGA)))
    for c in range(1, len(CABECALHO_SGA) + 1):
        ws.cell(row=2, column=c).number_format = "@"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO_SGA))}2"

    ws = wb.create_sheet(ABA_INTEGRACAO)
    cabecalho(ws, CABECALHO_INTEGRACAO, calculadas=range(G.N_COLUNAS, len(CABECALHO_INTEGRACAO)))
    for c in range(1, len(CABECALHO_INTEGRACAO) + 1):
        cel = ws.cell(row=2, column=c)
        cel.number_format = "DD/MM/YYYY" if c - 1 in G.COLUNAS_DATA else "0" if c - 1 == G.DURACAO else "@"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO_INTEGRACAO))}2"

    ws = wb.create_sheet(ABA_SIGA)
    nomes = list(CABECALHO_SIGA)
    if siga_valores:
        for c, v in enumerate(siga_valores[0][:I.SIGA_COLUNAS]):
            if texto(v).strip():
                nomes[c] = texto(v).strip()
    calculadas = set(I.SIGA_CALCULADAS) | set(I.SIGA_FORMULAS)
    cabecalho(ws, nomes, calculadas=calculadas)
    destino_linha = 2
    for _, linha in linhas_siga(siga_valores):
        for c in range(I.SIGA_COLUNAS):
            valor = I.SIGA_FORMULAS[c].format(r=destino_linha) if c in I.SIGA_FORMULAS else _valor_modelo(linha[c])
            if valor is None or valor == "":
                continue
            cel = ws.cell(row=destino_linha, column=c + 1, value=valor)
            if isinstance(valor, (datetime.date, datetime.datetime)):
                cel.number_format = "DD/MM/YYYY"
        destino_linha += 1
    ultima = max(destino_linha - 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(I.SIGA_COLUNAS)}{ultima}"

    # Mesmos alertas visuais que a planilha original tinha.
    alerta = PatternFill("solid", bgColor="F8CBAD")
    aviso = PatternFill("solid", bgColor="FFE699")
    ws.conditional_formatting.add("AE2:AG20000", FormulaRule(formula=["$AE2<>$AG2"], fill=alerta))
    for faixa in ("AJ2:AJ20000", "AQ2:AQ20000"):
        primeira = faixa.split(":")[0]
        ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'{primeira}="SETOR DESCONHECIDO"'], fill=alerta))
        ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'{primeira}="SEM TAREFA"'], fill=aviso))
    ws.conditional_formatting.add("AH2:AH20000", FormulaRule(fill=alerta, formula=[
        'AND(OR(UPPER(TRIM($AE2))="NEGADO",UPPER(TRIM($AG2))="NEGADO"),'
        'OR(TRIM($AH2)="",UPPER(TRIM($AH2))="NÃO INFORMADO",UPPER(TRIM($AH2))="NAO INFORMADO"))']))

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


# ----------------------------------------------------------------------------- escrita por dentro do .xlsx

def _xml(s: str) -> str:
    s = _ILEGAL.sub("", s.replace("\r\n", "\n").replace("\r", "\n"))
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _celula(ref: str, valor, estilo: str | None) -> str:
    s = f' s="{estilo}"' if estilo else ""
    if valor is None or valor == "":
        return f'<c r="{ref}"{s}/>' if estilo else ""
    if isinstance(valor, Formula):
        return f'<c r="{ref}"{s}><f>{_xml(valor.texto)}</f></c>'
    if isinstance(valor, bool):
        return f'<c r="{ref}"{s} t="b"><v>{int(valor)}</v></c>'
    if isinstance(valor, (int, float)):
        return f'<c r="{ref}"{s}><v>{valor!r}</v></c>'
    if isinstance(valor, datetime.datetime):
        if valor.time() != datetime.time(0):  # data e hora: a parte decimal é a fração do dia
            fracao = (valor.hour * 3600 + valor.minute * 60 + valor.second) / 86400
            return f'<c r="{ref}"{s}><v>{(valor.date() - _EPOCA).days + fracao!r}</v></c>'
        valor = valor.date()
    if isinstance(valor, datetime.date):
        return f'<c r="{ref}"{s}><v>{(valor - _EPOCA).days}</v></c>'
    t = _xml(str(valor))
    preservar = ' xml:space="preserve"' if t != t.strip() or "\n" in t else ""
    return f'<c r="{ref}"{s} t="inlineStr"><is><t{preservar}>{t}</t></is></c>'


def _limites_dados(xml: str) -> tuple[str, int, int]:
    """Devolve (xml, início, fim) do conteúdo de <sheetData>."""
    if "<sheetData/>" in xml:
        xml = xml.replace("<sheetData/>", "<sheetData></sheetData>", 1)
    ini = xml.find("<sheetData>")
    fim = xml.rfind("</sheetData>")
    if ini < 0 or fim < 0:
        raise ValueError("Aba sem <sheetData>: arquivo .xlsx inesperado.")
    return xml, ini + len("<sheetData>"), fim


def _linhas_xml(corpo: str):
    """Itera (nº, atributos, conteúdo, xml original) das linhas de um <sheetData>."""
    contador = 0
    for m in _LINHA.finditer(corpo):
        atributos = m.group(1) if m.group(1) is not None else m.group(2)
        r = re.search(r'\br="(\d+)"', atributos)
        contador = int(r.group(1)) if r else contador + 1
        yield contador, atributos, m.group(3) or "", m.group(0)


def _celulas_xml(conteudo: str):
    """Itera (índice da coluna, estilo, tem conteúdo?, xml original) das células de uma linha."""
    for m in _CELULA.finditer(conteudo):
        atributos = m.group(1) if m.group(1) is not None else m.group(2)
        ref = _REF.search(atributos)
        if not ref:
            continue
        estilo = _ESTILO.search(atributos)
        yield (column_index_from_string(ref.group(1)) - 1, estilo.group(1) if estilo else None,
               bool(_CONTEUDO.search(m.group(3) or "")), m.group(0))


def substituir_dados(xml: str, linhas, n_colunas: int) -> tuple[str, int]:
    """Troca todas as linhas de dados de uma aba, mantendo o cabeçalho (linha 1) e o estilo das colunas.

    Devolve (novo xml, nº da última linha).
    """
    xml, ini, fim = _limites_dados(xml)
    corpo = xml[ini:fim]
    cabecalho = ""
    estilos: list[collections.Counter] = [collections.Counter() for _ in range(n_colunas)]
    for numero, _, conteudo, original in _linhas_xml(corpo[:2_000_000]):
        if numero == 1:
            cabecalho = original
        elif numero <= 31:
            for coluna, estilo, _, _ in _celulas_xml(conteudo):
                if estilo and coluna < n_colunas:
                    estilos[coluna][estilo] += 1
        else:
            break
    estilo_coluna = [e.most_common(1)[0][0] if e else None for e in estilos]

    partes = [cabecalho]
    numero = 1
    for linha in linhas:
        numero += 1
        # A linha 2 leva todas as células (mesmo vazias) para continuar servindo de molde de estilo.
        celulas = [_celula(f"{_LETRAS[c]}{numero}", linha[c], estilo_coluna[c])
                   for c in range(min(n_colunas, len(linha))) if numero == 2 or (linha[c] is not None and linha[c] != "")]
        partes.append(f'<row r="{numero}">{"".join(celulas)}</row>')
    if numero == 1:  # sem dados: mantém a linha-molde vazia
        partes.append('<row r="2">' + "".join(_celula(f"{_LETRAS[c]}2", None, estilo_coluna[c]) for c in range(n_colunas)) + "</row>")
    ultima = max(numero, 2)
    xml = xml[:ini] + "".join(partes) + xml[fim:]

    faixa = f"A1:{_LETRAS[n_colunas - 1]}{ultima}"
    xml = re.sub(r'<dimension ref="[^"]*"\s*/>', f'<dimension ref="{faixa}"/>', xml, count=1)
    # O filtro continua na linha de cabeçalho, mas os critérios ativos são limpos (as linhas mudaram).
    xml = re.sub(r"<autoFilter\b[^>]*?/>|<autoFilter\b[^>]*>.*?</autoFilter>", f'<autoFilter ref="{faixa}"/>', xml, count=1, flags=re.S)
    return xml, ultima


def atualizar_colunas(xml: str, alteracoes: dict[int, dict[int, object]]) -> str:
    """Grava valores em células específicas, preservando todas as outras células e linhas."""
    xml, ini, fim = _limites_dados(xml)
    linhas: dict[int, tuple[str, dict[int, tuple[str | None, bool, str]]]] = {}
    estilos: dict[int, collections.Counter] = collections.defaultdict(collections.Counter)
    for numero, atributos, conteudo, _ in _linhas_xml(xml[ini:fim]):
        celulas = {}
        for coluna, estilo, tem_conteudo, original in _celulas_xml(conteudo):
            celulas[coluna] = (estilo, tem_conteudo, original)
            if estilo and numero > 1:
                estilos[coluna][estilo] += 1
        linhas[numero] = (atributos, celulas)
    estilo_padrao = {c: e.most_common(1)[0][0] for c, e in estilos.items()}

    for numero, colunas in alteracoes.items():
        atributos, celulas = linhas.get(numero) or (f' r="{numero}"', {})
        for coluna, valor in colunas.items():
            estilo, tem_conteudo, _ = celulas.get(coluna, (None, False, ""))
            if isinstance(valor, SeVazia):
                if tem_conteudo:
                    continue
                valor = valor.valor
            estilo = estilo or estilo_padrao.get(coluna)
            novo = _celula(f"{_LETRAS[coluna]}{numero}", valor, estilo)
            if novo:
                celulas[coluna] = (estilo, True, novo)
            else:
                celulas.pop(coluna, None)
        linhas[numero] = (atributos, celulas)

    partes = []
    for numero in sorted(linhas):
        atributos, celulas = linhas[numero]
        atributos = re.sub(r'\s+spans="[^"]*"', "", atributos)  # dica opcional que pode ficar desatualizada
        conteudo = "".join(celulas[c][2] for c in sorted(celulas))
        partes.append(f"<row{atributos}>{conteudo}</row>")
    return xml[:ini] + "".join(partes) + xml[fim:]


class EditorXlsx:
    """Abre um .xlsx como pacote zip e regrava apenas as partes alteradas."""

    def __init__(self, caminho: Path):
        self.caminho = Path(caminho)
        with zipfile.ZipFile(self.caminho) as z:
            self._workbook = z.read("xl/workbook.xml").decode("utf-8")
            relacoes = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        alvos = {}
        for tag in re.findall(r"<Relationship\b[^>]*>", relacoes):
            a = dict(_ATRIBUTO.findall(tag))
            alvo = a.get("Target", "")
            alvos[a.get("Id")] = alvo.lstrip("/") if alvo.startswith("/") else "xl/" + alvo
        self.abas: dict[str, str] = {}      # nome da aba -> parte do zip
        self.ordem: list[str] = []
        for tag in re.findall(r"<sheet\b[^>]*>", self._workbook):
            a = dict(_ATRIBUTO.findall(tag))
            rid = a.get("r:id") or next((v for k, v in a.items() if k.endswith(":id")), None)
            nome = html.unescape(a.get("name", ""))
            if rid in alvos:
                self.abas[nome] = alvos[rid]
                self.ordem.append(nome)
        self._novas: dict[str, str] = {}

    def nome_real(self, aba: str) -> str:
        real = _achar_aba(self.ordem, aba)
        if real is None:
            raise KeyError(f"A planilha {self.caminho.name} não tem a aba {aba}.")
        return real

    def xml(self, aba: str) -> str:
        parte = self.abas[self.nome_real(aba)]
        if parte in self._novas:
            return self._novas[parte]
        with zipfile.ZipFile(self.caminho) as z:
            return z.read(parte).decode("utf-8")

    def definir_xml(self, aba: str, xml: str) -> None:
        self._novas[self.abas[self.nome_real(aba)]] = xml

    def substituir_aba(self, aba: str, linhas, n_colunas: int) -> int:
        xml, ultima = substituir_dados(self.xml(aba), linhas, n_colunas)
        self.definir_xml(aba, xml)
        self._ajustar_filtro(self.nome_real(aba), n_colunas, ultima)
        return ultima

    def atualizar_aba(self, aba: str, alteracoes: dict[int, dict[int, object]]) -> None:
        self.definir_xml(aba, atualizar_colunas(self.xml(aba), alteracoes))

    def _ajustar_filtro(self, nome: str, n_colunas: int, ultima: int) -> None:
        """Mantém o nome interno do filtro automático (_FilterDatabase) coerente com a aba."""
        indice = str(self.ordem.index(nome))
        faixa = "'{}'!$A$1:${}${}".format(nome.replace("'", "''"), _LETRAS[n_colunas - 1], ultima)

        def trocar(m):
            a = dict(_ATRIBUTO.findall(m.group(1)))
            if a.get("name") == "_xlnm._FilterDatabase" and a.get("localSheetId") == indice:
                return f"<definedName{m.group(1)}>{_xml(faixa)}</definedName>"
            return m.group(0)

        self._workbook = re.sub(r"<definedName\b([^>]*)>([^<]*)</definedName>", trocar, self._workbook)

    def salvar(self, destino: Path) -> None:
        workbook = self._workbook
        m = re.search(r"<calcPr\b[^>]*/>", workbook)
        if m:  # manda o Excel recalcular as fórmulas ao abrir
            tag = re.sub(r'\s*fullCalcOnLoad="[^"]*"', "", m.group(0))
            workbook = workbook.replace(m.group(0), tag[:-2].rstrip() + ' fullCalcOnLoad="1"/>', 1)
        else:
            for ancora in ("</definedNames>", "</externalReferences>", "</functionGroups>", "</sheets>"):
                if ancora in workbook:
                    workbook = workbook.replace(ancora, ancora + '<calcPr fullCalcOnLoad="1"/>', 1)
                    break
        novas = dict(self._novas)
        novas["xl/workbook.xml"] = workbook
        with zipfile.ZipFile(self.caminho) as origem, \
                zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as saida:
            for info in origem.infolist():
                nome = info.filename
                if info.is_dir() or nome == "xl/calcChain.xml":
                    continue  # a cadeia de cálculo antiga não vale mais; o Excel recria
                if nome in novas:
                    dados = novas[nome].encode("utf-8")
                else:
                    dados = origem.read(nome)
                    if nome == "[Content_Types].xml":
                        dados = re.sub(rb'<Override\b[^>]*PartName="/xl/calcChain\.xml"[^>]*/>', b"", dados)
                    elif nome == "xl/_rels/workbook.xml.rels":
                        dados = re.sub(rb'<Relationship\b[^>]*Target="calcChain\.xml"[^>]*/>', b"", dados)
                saida.writestr(nome, dados)
        # Confere se o que foi escrito é um zip íntegro com XML bem formado.
        with zipfile.ZipFile(destino) as z:
            if z.testzip() is not None:
                raise ValueError("Falha ao gravar a planilha (zip corrompido).")
            for nome in novas:
                expat.ParserCreate().Parse(z.read(nome), True)


# ----------------------------------------------------------------------------- gravação do resultado

def planilha_bloqueada(caminho: Path) -> bool:
    try:
        with open(caminho, "r+b"):
            return False
    except PermissionError:
        return True


def _linha_integracao(linha: list, extra: list[str]) -> list:
    saida = list(linha)
    for c in G.COLUNAS_DATA:
        if isinstance(saida[c], str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", saida[c]):
            saida[c] = datetime.date.fromisoformat(saida[c])
    return saida + extra


def gravar(planilha: Path, sga: list[list[str]], integ: list[list], resultado: I.Resultado,
           pasta_backups: Path, backups_manter: int, fazer_backup: bool = True, siga_nova: list[list] | None = None) -> dict:
    """Grava as abas SGA e INTEGRAÇÃO e as colunas calculadas da aba SIGA. Faz backup antes.
    `siga_nova`: linhas (A..AS) do relatório de tarefas do SIGA, que substituem as da aba antes do cálculo."""
    if planilha_bloqueada(planilha):
        raise PlanilhaAberta(f"A planilha {planilha} está aberta em outro programa. Feche-a e rode de novo.")
    editor = EditorXlsx(planilha)
    if siga_nova is not None:
        editor.substituir_aba(ABA_SIGA, siga_nova, I.SIGA_COLUNAS)
    ultima_sga = editor.substituir_aba(ABA_SGA, (l + e for l, e in zip(sga, resultado.sga_extra)), len(CABECALHO_SGA))
    ultima_integ = editor.substituir_aba(
        ABA_INTEGRACAO, (_linha_integracao(l, e) for l, e in zip(integ, resultado.integ_extra)), len(CABECALHO_INTEGRACAO))
    alteracoes: dict[int, dict[int, object]] = {}
    for numero, valores in resultado.siga.items():
        colunas: dict[int, object] = dict(valores)
        for coluna, formula in I.SIGA_FORMULAS.items():
            colunas[coluna] = SeVazia(Formula(formula.format(r=numero)))
        alteracoes[numero] = colunas
    editor.atualizar_aba(ABA_SIGA, alteracoes)
    backup = salvar_editor(editor, pasta_backups, backups_manter, fazer_backup)
    return {"sga_ultima_linha": ultima_sga, "integracao_ultima_linha": ultima_integ, "backup": backup}


def salvar_editor(editor: EditorXlsx, pasta_backups: Path, backups_manter: int, fazer_backup: bool = True) -> str:
    """Grava as alterações do editor no próprio arquivo, com cópia de segurança da versão anterior.

    Devolve o nome do backup ('' se não foi feito).
    """
    planilha = editor.caminho
    temporario = planilha.with_name(f"~gravando_{planilha.name}")
    copia = None
    try:
        editor.salvar(temporario)
        if fazer_backup:
            pasta_backups.mkdir(parents=True, exist_ok=True)
            copia = pasta_backups / f"{planilha.stem}_{datetime.datetime.now():%Y%m%d_%H%M%S}{planilha.suffix}"
            shutil.copy2(planilha, copia)
        try:
            temporario.replace(planilha)
        except PermissionError as e:
            raise PlanilhaAberta(f"A planilha {planilha} está aberta em outro programa. Feche-a e rode de novo.") from e
    finally:
        temporario.unlink(missing_ok=True)
    antigos = sorted(pasta_backups.glob(f"{planilha.stem}_*{planilha.suffix}")) if pasta_backups.exists() else []
    for velho in antigos[:-backups_manter] if backups_manter > 0 else []:
        velho.unlink(missing_ok=True)
    return copia.name if copia else ""
