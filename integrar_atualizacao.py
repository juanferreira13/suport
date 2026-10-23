"""Orquestrador seguro: SIGA v3 -> valida 7 XLSX -> ENTRADA -> pipeline original.
Execute pelo 'Atualizar agora.bat' na raiz de grupo support.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
V3 = ROOT / 'SIGA_v3_exportacao_final' / 'exportar_v3.py'
PROCESSADOR = ROOT / 'sistema' / 'rodar.bat'
ENTRADA = ROOT / 'ENTRADA'
RESULTADOS = ROOT / 'RESULTADOS'
LOGS = ROOT / 'LOGS_SIGA_V3'
LOCK = ROOT / '.siga_v3_atualizacao.lock'
ENVIAR = ROOT / 'enviar_para_vercel.py'  # envio opcional dos paineis para o site (Vercel)


RE_NUMERO = re.compile(r'^(\d{4})-\d+$')


def checar_arquivo(caminho, tipo, ano=None):
    """Mesma regra de exportar_v3.conferir_ano (manter as duas iguais).

    O Excel do Dashboard traz linhas de grupo (ex.: 'COLISAO') e eventos com
    numero de anos anteriores cadastrados no ano; 'max_row' vem None nesses
    arquivos (read_only), por isso as linhas sao contadas uma a uma.
    """
    if not caminho.is_file() or caminho.stat().st_size < 500:
        raise RuntimeError(f'XLSX ausente ou muito pequeno: {caminho}')
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = wb.active.iter_rows(values_only=True)
        next(linhas, None)  # titulo
        cab = list(next(linhas, ()))
        dados = [row for row in linhas if any(v is not None for v in row)]
    finally:
        wb.close()
    if not dados:
        raise RuntimeError(f'Planilha sem registros: {caminho.name}')
    if tipo != 'total':
        return len(dados)
    indices = [i for i, val in enumerate(cab) if val and 'mero' in str(val).lower()]
    if not indices:
        raise RuntimeError(f'Cabecalho Numero nao encontrado: {caminho.name}')
    idx = indices[0]
    anos = [int(m.group(1)) for row in dados
            if idx < len(row) and (m := RE_NUMERO.match(str(row[idx] or '').strip()))]
    if not anos:
        raise RuntimeError(f'Arquivo sem eventos: {caminho.name}')
    futuros = sorted({a for a in anos if a > ano})
    if futuros:
        raise RuntimeError(f'Ano divergente em {caminho.name}: eventos de {futuros}')
    if sum(a == ano for a in anos) < 0.8 * len(anos):
        raise RuntimeError(f'Ano divergente em {caminho.name}: maioria nao e de {ano}')
    return len(anos)


def validar_manifesto(log, pasta):
    if not isinstance(log, dict):
        raise RuntimeError('JSON de exportacao invalido')
    anos = log.get('anos') or []
    esperados = list(range(datetime.now().year, 2020, -1))
    if len(anos) != len(esperados) or [x.get('ano') for x in anos] != esperados:
        raise RuntimeError('A lista de anos da exportacao nao confere')
    arquivos = []
    for item, ano in zip(anos, esperados):
        if item.get('status') != 'OK':
            raise RuntimeError(f'Exportacao falhou para {ano}: {item.get("erro", "sem detalhes")}')
        origem = Path(item.get('arquivo', ''))
        if origem.parent.resolve() != pasta.resolve() or not origem.name.startswith(f'Total_Eventos_Cadastrados_{ano}_'):
            raise RuntimeError(f'Caminho invalido para {ano}: {origem}')
        contagem = checar_arquivo(origem, 'total', ano)
        if contagem != item.get('registros'):
            raise RuntimeError(f'Contagem do Excel {ano}: {contagem} != {item.get("registros")}')
        arquivos.append(origem)
    ctrl = log.get('controle') or {}
    if ctrl.get('status') != 'OK':
        raise RuntimeError('Falha no Controle de Eventos')
    origem = Path(ctrl.get('arquivo', ''))
    if origem.parent.resolve() != pasta.resolve() or not origem.name.startswith('Controle_Eventos_'):
        raise RuntimeError(f'Caminho invalido para Controle: {origem}')
    checar_arquivo(origem, 'controle')
    arquivos.append(origem)
    if len(set(map(str, arquivos))) != 7:
        raise RuntimeError('Arquivos repetidos no manifesto')
    return arquivos


CODIGO_ERRO_LOGIN = 3  # mesmo codigo de exportar_v3.py


def progresso(pct, msg):
    """Linha lida pela tela de atualizacao (atualizar_tela.py) para a barra de progresso."""
    print(f'@@PROGRESSO {int(pct)} {msg}', flush=True)


def processo_vivo(lock):
    """O processo que criou o bloqueio ainda existe? (bloqueio fica se a janela for fechada a forca)"""
    try:
        pid = int(lock.read_text(encoding='utf-8').split()[0].split('=')[1])
    except (OSError, ValueError, IndexError):
        return False
    import ctypes
    k32 = ctypes.windll.kernel32
    h = k32.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
    if not h:
        return False
    try:
        codigo = ctypes.c_ulong()
        return bool(k32.GetExitCodeProcess(h, ctypes.byref(codigo))) and codigo.value == 259  # STILL_ACTIVE
    finally:
        k32.CloseHandle(h)


def detalhe(msg):
    """Sub-etapa mostrada na tela de atualizacao (ex.: 'API do SGA: 1.234 itens carregados')."""
    print(f'@@DETALHE {msg}', flush=True)


# Mensagens do processamento (sistema/app/sinistro/pipeline.py) que viram etapa na tela:
# (trecho da linha, %, etapa). A linha inteira tambem aparece como detalhe.
ETAPAS_PROCESSAMENTO = [
    ('Primeira carga', 88, 'Carregando o histórico das planilhas'),
    ('Relatório SGA', 89, 'Carregando o relatório do SGA'),
    ('API SGA: autenticado', 89, 'Carregando eventos pela API do SGA'),
    ('API SGA não configurada', 89, 'Carregando o relatório do SGA'),
    ('Google: lendo', 90, 'Carregando as planilhas do Google'),
    ('DashBoard SIGA', 91, 'Integrando os eventos do SIGA'),
    ('Relatório de tarefas do SIGA', 91, 'Integrando as tarefas do SIGA'),
    ('Aba SGA:', 92, 'Montando as abas SGA, INTEGRAÇÃO e SIGA'),
    ('Gravado', 93, 'Gravando as planilhas'),
    ('Página ', 94, 'Gerando os painéis HTML'),
]


RE_PREFIXO_LOG = re.compile(r'^\d\d/\d\d/\d{4} \d\d:\d\d:\d\d \w+\s+')


def processar(comando):
    """Roda o processamento repassando a saida e traduzindo as mensagens em etapas da tela."""
    proc = subprocess.Popen(comando, cwd=str(ROOT), stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    for bruta in proc.stdout:
        linha = bruta.decode('utf-8', errors='replace').rstrip()
        print(linha, flush=True)
        for trecho, pct, etapa in ETAPAS_PROCESSAMENTO:
            if trecho in linha:
                progresso(pct, etapa)
                break
        if 'API SGA' in linha or 'DashBoard SIGA' in linha or 'Google' in linha:
            # tira data/hora/nivel do inicio (formato de sinistro/util.py) para a tela
            detalhe(RE_PREFIXO_LOG.sub('', linha).strip()[:160])
    return proc.wait()


def copiar_resultados(destino):
    """Copia as planilhas Excel e os paineis HTML de RESULTADOS para a pasta escolhida na tela.

    RESULTADOS continua sendo a pasta de trabalho (o sistema le as colunas digitadas pela equipe
    de la); o destino recebe uma copia atualizada a cada execucao.
    """
    if destino.resolve() == RESULTADOS.resolve():
        return
    progresso(95, f'Salvando planilhas e painéis em {destino}')
    destino.mkdir(parents=True, exist_ok=True)
    falhas = []
    for origem in sorted(RESULTADOS.iterdir()):
        if origem.suffix.lower() not in ('.xlsx', '.html') or origem.name.startswith('~$'):
            continue
        try:
            shutil.copy2(origem, destino / origem.name)
            print(f'  copiado: {origem.name}', flush=True)
        except OSError as e:
            falhas.append(f'{origem.name} ({e.strerror or e})')
    if falhas:
        raise RuntimeError('Atualizacao feita, mas nao consegui gravar em ' + str(destino) + ': '
                           + '; '.join(falhas) + '. Feche esses arquivos e rode de novo.')


def enviar_para_site():
    """Envia os paineis gerados para o site (Vercel), se ele estiver configurado.

    So roda quando VERCEL_SITE_URL e VERCEL_UPLOAD_SECRET existem (em sistema\\configuracao.env
    ou no ambiente). Uma falha aqui NAO derruba a atualizacao: os arquivos locais ja foram
    gerados e salvos; o envio ao site e um extra.
    """
    if not ENVIAR.is_file():
        return
    try:
        if str(ROOT) not in sys.path:
            sys.path.insert(0, str(ROOT))
        import enviar_para_vercel as site
    except Exception as e:  # noqa: BLE001 - nunca interromper a atualizacao por causa do envio
        print(f'Aviso: nao consegui carregar o envio para o site: {e}', flush=True)
        return
    env = site.ler_env(site.CONFIG)
    url = (os.environ.get('VERCEL_SITE_URL') or env.get('VERCEL_SITE_URL') or '').strip()
    segredo = (os.environ.get('VERCEL_UPLOAD_SECRET') or env.get('VERCEL_UPLOAD_SECRET') or '').strip()
    if not url or not segredo:
        return  # site nao configurado: nao faz nada (sem ruido)
    # So envia os paineis que realmente existem (evita aviso de arquivo faltando).
    nomes = [nome for nome in site.PAINEIS if site.arquivo_do_painel(nome, env).is_file()]
    if not nomes:
        return
    progresso(97, 'Enviando os painéis para o site')
    detalhe('Enviando os painéis para o site (Vercel)')
    try:
        codigo = site.main(nomes)
        detalhe('Painéis enviados para o site' if codigo == 0
                else 'Não consegui enviar os painéis para o site (a atualização local foi concluída)')
    except Exception as e:  # noqa: BLE001 - idem
        detalhe(f'Falha ao enviar para o site: {e} (a atualização local foi concluída)')


def main():
    if not V3.is_file() or not PROCESSADOR.is_file():
        raise RuntimeError('Nao encontrei a v3 ou sistema\\rodar.bat')
    LOGS.mkdir(exist_ok=True)
    if LOCK.exists() and not processo_vivo(LOCK):
        print('Bloqueio antigo (atualizacao interrompida) removido.', flush=True)
        LOCK.unlink(missing_ok=True)
    if LOCK.exists():
        raise RuntimeError('Existe outra atualizacao em andamento ou um bloqueio antigo: ' + str(LOCK))
    LOCK.write_text(f'pid={os.getpid()} data={datetime.now().isoformat()}\n', encoding='utf-8')
    try:
        pasta = Path(tempfile.mkdtemp(prefix='siga_', dir=ROOT))
        log = pasta / 'resultado_exportacao_v3.json'
        env = os.environ.copy()
        env['SIGA_V3_SAIDA'] = str(pasta)
        env['SIGA_V3_LOG'] = str(log)
        # Sem janela do SIGA quando ninguem definiu (ex.: tarefa agendada ou modo --texto).
        # A tela de atualizacao ja define SIGA_V3_OCULTO=1, entao aqui nao muda nada por ela.
        env.setdefault('SIGA_V3_OCULTO', '1')
        progresso(1, 'Preparando')
        print('[1/3] Exportacao SIGA v3 em pasta isolada...', flush=True)
        ret = subprocess.call([sys.executable, str(V3)], cwd=str(V3.parent), env=env)
        if ret == CODIGO_ERRO_LOGIN:
            shutil.rmtree(pasta, ignore_errors=True)
            print('ERRO: login ou senha do SIGA recusados.', flush=True)
            return CODIGO_ERRO_LOGIN
        if ret != 0:
            raise RuntimeError(f'Exportador v3 retornou erro {ret}. Arquivos nao transferidos. Veja {log}')
        if not log.is_file():
            raise RuntimeError('Exportacao nao criou JSON de diagnostico')
        dados = json.loads(log.read_text(encoding='utf-8'))
        progresso(82, 'Conferindo os arquivos baixados')
        print('[2/3] Validando 6 anos e Controle...', flush=True)
        arquivos = validar_manifesto(dados, pasta)
        ENTRADA.mkdir(exist_ok=True)
        # Prepara os 7 arquivos com nomes novos, sem tocar em dados existentes.
        for origem in arquivos:
            destino = ENTRADA / origem.name
            if destino.exists():
                raise RuntimeError(f'Arquivo de destino ja existe: {destino}')
        transferidos = []
        try:
            for origem in arquivos:
                destino = ENTRADA / origem.name
                shutil.copy2(origem, destino)
                transferidos.append(destino)
        except Exception:
            for f in transferidos:
                f.unlink(missing_ok=True)
            raise
        print(f'  {len(transferidos)} relatórios novos em ENTRADA.', flush=True)
        progresso(88, 'Atualizando planilhas e painéis (pode levar alguns minutos)')
        print('[3/3] Atualizando planilhas e paineis Grupo Support...', flush=True)
        ret = processar(['cmd', '/c', str(PROCESSADOR), 'executar.py', *sys.argv[1:]])
        shutil.copy2(log, LOGS / f'resultado_exportacao_v3_{datetime.now():%Y%m%d_%H%M%S}.json')
        if ret != 0:
            raise RuntimeError(f'Processamento do Grupo Support falhou (codigo {ret}); confira ENTRADA e logs')
        shutil.rmtree(pasta, ignore_errors=True)  # arquivos ja estao em ENTRADA e o JSON em LOGS
        if os.environ.get('SIGA_V3_DESTINO'):
            copiar_resultados(Path(os.environ['SIGA_V3_DESTINO']))
        enviar_para_site()
        progresso(100, 'Atualização concluída')
        print('ATUALIZACAO CONCLUIDA COM SUCESSO.', flush=True)
        return 0
    finally:
        LOCK.unlink(missing_ok=True)


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as exc:
        print('ERRO:', exc, file=sys.stderr, flush=True)
        sys.exit(1)
