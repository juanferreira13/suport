"""SIGA v3: 6 relatorios pela API; Controle de Eventos pela interface da v2.
Nao modifica arquivos da v2. Valida ano e total antes de publicar cada XLSX.
"""
import base64
import json
import os
import re
import sys
import time
import traceback
from datetime import datetime, date
import shutil
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
CANDIDATOS = [ROOT / 'siga_automacao-v2' / 'siga_automacao', ROOT / 'siga_automacao']
PROJETO = next((p for p in CANDIDATOS if (p/'siga.py').is_file() and (p/'relatorios.py').is_file()), None)
if not PROJETO:
    raise SystemExit('Nao encontrei a v2. Extraia SIGA_v3_exportacao dentro da pasta grupo support.')
sys.path.insert(0, str(PROJETO))
import siga, relatorios
import credenciais_portatil as credenciais

SAIDA = Path(os.environ['SIGA_V3_SAIDA']) if os.environ.get('SIGA_V3_SAIDA') else ROOT / 'SAIDA_SIGA_V3_TESTE'
LOG = Path(os.environ['SIGA_V3_LOG']) if os.environ.get('SIGA_V3_LOG') else ROOT / 'resultado_exportacao_v3.json'
URL = 'https://siga.arcoconsultoria.com/SIGA/wsTotalEventosCadastrado.rule?sys=SIG'
COLUNAS = [
 ('Numero','Numero'),('Integração','Integracao'),('Origem','Origem'),('Regional','regional'),
 ('Cooperativa','cooperativa'),('Situação Evento','Situacao_Evento'),('Proprietário','Proprietario'),
 ('Estado','Estado'),('Tipo Veic.','Tipo_Veic'),('Cobertura','Cobertura'),('Placa','Placa'),
 ('Chassi','Chassi'),('Cadastro','Cadastro'),('Abertura','Abertura'),('Entrada Oficina','Entrada_Oficina'),
 ('Liberação','Liberacao'),('Regulador','Regulador'),('Duração','Duracao'),('Oficina','Oficina'),
 ('Peças Oficina','Pecas_Oficina'),('Mão de Obra','Mao_de_Obra'),('Peças','Pecas'),
 ('Reembolso','Reembolso'),('Carro Reserva','Carro_Reserva'),('Diversos','Diversos'),
 ('Participação','Participacao'),('Indenização','Indenizacao'),('Total Rateio','Total_Rateio')]


RE_NUMERO = re.compile(r'^(\d{4})-\d+$')


def conferir_ano(numeros, ano):
    """Confere se os numeros de evento (ex.: 2021-00090) sao do ano pedido.

    O Excel oficial traz linhas de grupo (ex.: 'COLISAO') e eventos com numero
    de anos anteriores cadastrados no ano (ex.: 77 eventos 2020- em 2021).
    Por isso exige: nenhum evento de ano posterior e a grande maioria do ano.
    Retorna a quantidade de eventos.
    """
    anos = [int(m.group(1)) for n in numeros if (m := RE_NUMERO.match(str(n or '').strip()))]
    if not anos:
        raise RuntimeError('Excel sem registros')
    futuros = sorted({a for a in anos if a > ano})
    if futuros:
        raise RuntimeError(f'Excel {ano} contem eventos de ano posterior: {futuros}')
    proprios = sum(a == ano for a in anos)
    if proprios < 0.8 * len(anos):
        raise RuntimeError(f'So {proprios} de {len(anos)} eventos sao de {ano}')
    return len(anos)


def contar_eventos_xlsx(caminho, ano):
    """Le a coluna Numero (linha 2 = cabecalho) e confere o ano. Fecha o arquivo antes de retornar."""
    from openpyxl import load_workbook
    book = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = book.active.iter_rows(values_only=True)
        next(linhas, None)  # titulo 'DashBoard SIGA'
        cabecalhos = next(linhas, ())
        indice = next((i for i, c in enumerate(cabecalhos) if c and 'mero' in str(c).lower()), None)
        if indice is None:
            raise RuntimeError('Cabecalho Numero nao encontrado no Excel')
        numeros = [linha[indice] for linha in linhas if indice < len(linha)]
    finally:
        book.close()
    return conferir_ano(numeros, ano)


def requisitar(dash, payload):
    # Requisicao a partir da pagina autenticada: usa cookies e contexto real do Dashboard.
    result = dash.evaluate('''async ({url, payload}) => {
      const r = await fetch(url, {method:'POST', credentials:'include',
        headers:{'Content-Type':'application/json'}, body:JSON.stringify(payload)});
      const texto = await r.text();
      let dados;
      try { dados = JSON.parse(texto); } catch(e) { throw new Error('Resposta nao e JSON. HTTP '+r.status+'; inicio: '+texto.slice(0,120)); }
      return {http:r.status, dados};
    }''', {'url': URL, 'payload': payload})
    if result['http'] != 200:
        raise RuntimeError(f"API HTTP {result['http']}")
    dados = result['dados']
    if not isinstance(dados, dict) or not isinstance(dados.get('tarefas'), list):
        raise RuntimeError('Resposta inesperada da API: tarefas ausentes')
    if str(dados.get('status','')).lower() not in ('sucess','success','sucesso'):
        raise RuntimeError('Status inesperado da API')
    return dados


def coletar(dash, ano, emp_id, mes):
    registros = []
    total = None
    tamanho = 250
    for pagina in range(2000):
        payload = {'emp_id':str(emp_id), 'ano':str(ano), 'mes':str(mes),
                   'relatorio':'relatorio', 'offset':len(registros), 'limit':tamanho}
        dados = requisitar(dash, payload)
        lote = dados['tarefas']
        if total is None:
            try: total = int(dados['recordTotal'])
            except (KeyError, ValueError, TypeError): raise RuntimeError('API nao informou recordTotal')
            if total <= 0: raise RuntimeError('API informou zero registros; nao salvarei arquivo vazio')
        elif int(dados.get('recordTotal', -1)) != total:
            raise RuntimeError('recordTotal mudou durante paginacao')
        if not lote: raise RuntimeError(f'API parou em {len(registros)} de {total} registros')
        if not all(isinstance(r,dict) for r in lote): raise RuntimeError('Registro inesperado')
        registros.extend(lote)
        print(f'  {ano}: {len(registros)}/{total}',flush=True)
        detalhe(f'carregando pela API do SIGA: {len(registros)} de {total} registros')
        if len(registros) >= total: break
    if len(registros) != total:
        raise RuntimeError(f'Quantidade inconsistente: {len(registros)} != {total}')
    # A coluna Numero e o identificador anual do SIGA, como 2021-00090.
    conferir_ano([r.get('Numero') for r in registros], ano)
    return registros


def salvar_excel(ano, registros):
    wb=Workbook();ws=wb.active;ws.title='Sheet1'
    ws.append(['DashBoard SIGA']);ws.append([nome for nome,_ in COLUNAS])
    for reg in registros:
        ws.append([reg.get(chave) for _,chave in COLUNAS])
    ws.freeze_panes='A3';ws.auto_filter.ref=f'A2:AB{ws.max_row}'
    for c in ws[2]:
        c.font=Font(color='FFFFFF',bold=True)
        c.fill=PatternFill('solid',fgColor='253858')
    ws.column_dimensions['A'].width=19
    ws.column_dimensions['B'].width=23
    ws.column_dimensions['G'].width=38
    ws.column_dimensions['S'].width=30
    SAIDA.mkdir(parents=True,exist_ok=True)
    caminho=SAIDA/f'Total_Eventos_Cadastrados_{ano}_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx'
    temp=caminho.with_suffix('.xlsx.tmp')
    wb.save(temp);temp.replace(caminho)
    return caminho


def fechar_relatorio(dash, timeout_ms=7000):
    """Fecha APENAS o popup de eventos, sem fechar a janela do Dashboard.

    Alguns layouts do SIGA nao usam `.modal.show`; por isso partimos da tabela
    real e procuramos seu container. Se houver falha, interrompemos a etapa
    seguinte para evitar que o filtro seja aplicado sobre o popup anterior.
    """
    if dash.is_closed():
        return
    tabela = dash.locator('#table_relatorio_eventos')
    if not tabela.count() or not tabela.first.is_visible():
        return

    # Primeiro procura o botao X no container da propria tabela (na imagem,
    # no canto superior direito), sem atingir botoes de fechar outras janelas.
    script = r"""() => {
      const t = document.querySelector('#table_relatorio_eventos');
      if (!t) return false;
      const vis = e => { const r=e.getBoundingClientRect();
        return r.width>0 && r.height>0 && getComputedStyle(e).visibility !== 'hidden'; };
      const caixas = [];
      let p=t.parentElement;
      for(let i=0;p&&i<10;p=p.parentElement,i++) {
        if(vis(p)) caixas.push(p);
        if(p.matches('.modal, [role="dialog"], .ui-dialog, .modal-content')) break;
      }
      for (const caixa of caixas) {
        const candidatos=[...caixa.querySelectorAll(
          'button, a, span, i, [role="button"], [data-dismiss], [data-bs-dismiss]')]
          .filter(vis).filter(e => {
            const a=(e.getAttribute('aria-label')||'').toLowerCase();
            const title=(e.getAttribute('title')||'').toLowerCase();
            const cls=String(e.className||'').toLowerCase();
            const txt=(e.textContent||'').trim();
            return /close|fechar/.test(a+' '+title+' '+cls) ||
              ['×','✕','✖','x'].includes(txt);
          });
        // Limita o alvo ao canto superior direito do popup para evitar
        // botoes "fechar" existentes dentro da tabela.
        const r=caixa.getBoundingClientRect();
        const alvos=candidatos.filter(e => { const x=e.getBoundingClientRect();
          return x.top < r.top+90 && x.right > r.right-110; });
        if(alvos.length) { alvos[0].click(); return true; }
      }
      return false;
    }"""
    dash.evaluate(script)
    try:
        tabela.first.wait_for(state='hidden', timeout=1800)
        return
    except Exception:
        pass

    # Fallback para modais Bootstrap.
    try:
        dash.keyboard.press('Escape')
        tabela.first.wait_for(state='hidden', timeout=1800)
        return
    except Exception:
        pass

    # Ultimo recurso: X visual do modal, sem coordenadas fixas da tela.
    # O retangulo escolhido precisa conter a tabela e ser parecido com popup.
    alvo=dash.evaluate("""() => {
      let e=document.querySelector('#table_relatorio_eventos');
      if(!e) return null;
      const vw=innerWidth, vh=innerHeight;
      while(e && e!==document.body) {
        const r=e.getBoundingClientRect();
        if(r.width>300 && r.width<vw && r.height>150 && r.height<vh &&
           r.left>=0 && r.top>=0 &&
           (e.matches('.modal-content,[role="dialog"],.ui-dialog') ||
            (getComputedStyle(e).position==='fixed' && r.width>vw*.5))) {
          return {x:r.right-22,y:r.top+19};
        }
        e=e.parentElement;
      }
      return null;
    }""")
    if alvo:
        dash.mouse.click(alvo['x'],alvo['y'])
        try:
            tabela.first.wait_for(state='hidden',timeout=1800)
            return
        except Exception:
            pass
    raise RuntimeError('Nao consegui fechar o popup do relatorio; interrompendo para evitar ano incorreto')


def preparar_ano(dash, ano):
    fechar_relatorio(dash)
    sel=dash.locator('#select_ano')
    dash.evaluate(relatorios.JS_MARCAR_BARRAS_VELHAS)
    sel.select_option(label=str(ano))
    dash.locator('button[title="Buscar"]').click()
    try:dash.wait_for_load_state('networkidle',timeout=45000)
    except Exception:pass
    dash.wait_for_timeout(2000)
    if sel.locator('option:checked').inner_text().strip()!=str(ano):
        raise RuntimeError(f'Dashboard alterou o ano solicitado {ano}: {sel.input_value()}')
    try:
        relatorios._esperar(dash,lambda:dash.evaluate(relatorios.JS_MARCAR_BARRA_TOTAL,False),25,'Grafico','Nao redesenhou')
    except relatorios.ErroPasso:
        relatorios._esperar(dash,lambda:dash.evaluate(relatorios.JS_MARCAR_BARRA_TOTAL,True),20,'Grafico','Nao encontrado')
    dash.locator('[data-siga-alvo]').scroll_into_view_if_needed()
    dash.wait_for_timeout(600)
    dash.evaluate(relatorios.JS_MARCAR_BARRA_TOTAL,True)
    dash.locator('[data-siga-alvo]').click()
    dash.locator('#table_relatorio_eventos').wait_for(state='visible',timeout=45000)
    dash.wait_for_timeout(800)
    if sel.locator('option:checked').inner_text().strip()!=str(ano):
        raise RuntimeError('Ano mudou ao abrir o popup')


def guardar(info):
    LOG.write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf8')


CODIGO_ERRO_LOGIN = 3  # integrar_atualizacao.py e atualizar_tela.py tratam este codigo


def detalhe(msg):
    """Sub-etapa da etapa atual (ex.: 'carregando pela API do SIGA: 250 de 4125'), mostrada na tela."""
    print(f'@@DETALHE {msg}', flush=True)


def progresso(pct, msg):
    """Linha lida pela tela de atualizacao (atualizar_tela.py) para a barra de progresso."""
    print(f'@@PROGRESSO {int(pct)} {msg}', flush=True)


def obter_credenciais():
    # A tela de atualizacao passa login/senha por variavel de ambiente (nada fica em disco
    # se 'Lembrar' estiver desmarcado). Sem tela: credencial salva ou pergunta no terminal.
    if os.environ.get('SIGA_V3_USUARIO') and os.environ.get('SIGA_V3_SENHA'):
        return os.environ['SIGA_V3_USUARIO'], os.environ['SIGA_V3_SENHA']
    creds=credenciais.carregar()
    if creds:
        return creds[:2]
    import getpass
    print('PRIMEIRO ACESSO SIGA: entre com suas credenciais (guardadas criptografadas neste Windows).', flush=True)
    usuario = input('Usuario SIGA: ').strip()
    senha = getpass.getpass('Senha SIGA: ')
    if not usuario or not senha: raise RuntimeError('Usuario ou senha vazios')
    credenciais.salvar(usuario, senha)
    return usuario, senha


class Previa:
    """Visao ao vivo do navegador oculto para a tela de atualizacao.

    Usa o espelhamento de tela do proprio Edge (Page.startScreencast) e grava o quadro mais
    recente em SIGA_V3_PREVIA; a tela (atualizar_tela.py) le esse arquivo. Sem a variavel, nao faz nada.
    """

    def __init__(self, ctx):
        caminho = os.environ.get('SIGA_V3_PREVIA')
        self.ctx, self.arquivo, self.cdp, self.ultimo = ctx, Path(caminho) if caminho else None, None, 0.0

    def seguir(self, pagina):
        """Passa a espelhar esta pagina (a principal ou a do Dashboard)."""
        if not self.arquivo:
            return
        if self.cdp:
            try:
                self.cdp.send('Page.stopScreencast'); self.cdp.detach()
            except Exception:
                pass
        try:
            cdp = self.ctx.new_cdp_session(pagina)
            cdp.on('Page.screencastFrame', lambda q: self._quadro(cdp, q))
            cdp.send('Page.startScreencast', {'format': 'jpeg', 'quality': 60, 'maxWidth': 720, 'maxHeight': 450})
            self.cdp = cdp
        except Exception as e:
            print('  Previa do navegador indisponivel:', e, flush=True)
            self.cdp = None

    def _quadro(self, cdp, quadro):
        try:
            cdp.send('Page.screencastFrameAck', {'sessionId': quadro['sessionId']})
        except Exception:
            return
        agora = time.time()
        if agora - self.ultimo < 0.4:  # a tela atualiza ~1x por segundo; nao precisa gravar todo quadro
            return
        self.ultimo = agora
        try:
            temp = self.arquivo.with_suffix('.tmp')
            temp.write_bytes(base64.b64decode(quadro['data']))
            os.replace(temp, self.arquivo)
        except OSError:
            pass  # a tela pode estar lendo o arquivo neste instante; vale o proximo quadro


def main():
    creds=obter_credenciais()
    info={'inicio':datetime.now().isoformat(timespec='seconds'),'projeto':str(PROJETO),'saida':str(SAIDA),'anos':[],'controle':None}
    guardar(info)
    oculto=os.environ.get('SIGA_V3_OCULTO')=='1'
    anos=list(range(date.today().year,2020,-1))
    progresso(3,'Abrindo navegador' + (' (oculto)' if oculto else ''))
    with sync_playwright() as pw:
        ctx=siga.abrir_navegador(pw,headless=oculto)
        try:
            page=ctx.pages[0] if ctx.pages else ctx.new_page()
            previa=Previa(ctx);previa.seguir(page)
            progresso(8,'Conectando ao SIGA')
            try:
                siga.entrar(page,*creds)
            except siga.ErroLogin as e:
                print(f'@@ERRO_LOGIN {e}',flush=True)
                info['login']={'status':'ERRO','erro':str(e)}
                info['fim']=datetime.now().isoformat(timespec='seconds');guardar(info)
                return CODIGO_ERRO_LOGIN
            print('Login OK');print('@@LOGIN_OK',flush=True)
            progresso(12,'Carregando o Dashboard Arco')
            dash=relatorios.abrir_dashboard_arco(ctx,page)
            previa.seguir(dash)
            dash.locator('li:text-is("Gerencial")').first.click()
            dash.locator('#select_ano').wait_for(state='visible',timeout=60000)
            for i,ano in enumerate(anos):
                progresso(15+i*60/len(anos),f'Carregando Total Eventos Cadastrados {ano} ({i+1} de {len(anos)})')
                print('\n---',ano,'---',flush=True)
                item={'ano':ano};info['anos'].append(item)
                try:
                    detalhe(f'selecionando {ano} no Dashboard')
                    preparar_ano(dash,ano)
                    # Usar o mesmo emp_id e mes observados na captura; se o contexto mudar, falha com zero.
                    # API pode responder 200 com recordTotal=0 mesmo quando a tabela tem dados.
                    # Nesse caso, baixar o Excel oficial pelo popup ja aberto.
                    try:
                        detalhe('consultando a API do SIGA')
                        registros=coletar(dash,ano,'141',str(datetime.now().month))
                        destino=salvar_excel(ano,registros)
                        # Conta do mesmo jeito que integrar_atualizacao.py confere depois.
                        item.update(status='OK',metodo='API',registros=contar_eventos_xlsx(destino,ano),arquivo=str(destino))
                    except Exception as erro_api:
                        print('  API nao validada:',erro_api)
                        print('  Usando exportacao oficial Excel do Dashboard...')
                        detalhe('baixando o Excel pelo Dashboard (download)')
                        # Garante que a tabela pertence ao ano solicitado.
                        selecionado=dash.locator('#select_ano option:checked').inner_text().strip()
                        if selecionado != str(ano):
                            raise RuntimeError(f'Ano no Dashboard e {selecionado}, esperado {ano}')
                        tabela=dash.locator('#table_relatorio_eventos')
                        tabela.wait_for(state='visible',timeout=60000)
                        try:
                            dash.locator('.dataTables_processing').first.wait_for(state='hidden',timeout=60000)
                        except Exception:
                            pass
                        botao=dash.locator('#table_relatorio_eventos_wrapper :is(button, a):has-text("Excel")').first
                        if not botao.count() or not botao.is_visible():
                            botao=dash.locator('.modal.show :is(button, a):has-text("Excel")').first
                        botao.wait_for(state='visible',timeout=30000)
                        with dash.expect_download(timeout=300000) as dl:
                            botao.click()
                        download=dl.value
                        SAIDA.mkdir(parents=True,exist_ok=True)
                        destino=SAIDA/f'Total_Eventos_Cadastrados_{ano}_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx'
                        download.save_as(destino)
                        # Confirma que o Excel e legivel e contem o ano correto.
                        # contar_eventos_xlsx fecha o arquivo antes; senao o unlink falha (WinError 32).
                        try:
                            n=contar_eventos_xlsx(destino,ano)
                        except Exception:
                            destino.unlink(missing_ok=True)
                            raise
                        item.update(status='OK',metodo='Excel oficial (fallback)',registros=n,arquivo=str(destino),aviso_api=str(erro_api))
                    print('OK',destino,'|',item['metodo'],'|',item['registros'],'registros')
                    detalhe(f"{item['registros']} eventos carregados")
                except Exception as e:
                    item.update(status='ERRO',erro=f'{type(e).__name__}: {e}')
                    print('FALHA',ano,item['erro'])
                finally:
                    guardar(info)
                    try:
                        fechar_relatorio(dash)
                        print('  Popup do relatorio fechado.')
                    except Exception as erro_fechar:
                        item['erro_fechamento']=str(erro_fechar)
                        item['status']='ERRO'
                        guardar(info)
                        print('  FALHA AO FECHAR POPUP:', erro_fechar)
                        print('  Interrompendo a sequencia para nao confundir os anos.')
                        break
            # Ao terminar os anos, fechar SOMENTE a aba/janela Gerencial.
            # A guia principal SIGA continua aberta no objeto `page`, como na v2.
            print('\n--- Encerrando guia Gerencial / Dashboard Arco ---', flush=True)
            try:
                if dash is page:
                    raise RuntimeError('Dashboard e pagina principal sao a mesma guia; fechamento bloqueado')
                if not dash.is_closed():
                    dash.close(run_before_unload=False)
                if not dash.is_closed():
                    raise RuntimeError('A guia Gerencial continua aberta')
                if page.is_closed():
                    raise RuntimeError('A guia principal do SIGA foi fechada inesperadamente')
                page.bring_to_front()
                previa.seguir(page)
                print('OK Guia Gerencial fechada. Voltando a pagina principal do SIGA.')
                info['fechamento_gerencial']={'status':'OK'}
            except Exception as e:
                info['fechamento_gerencial']={'status':'ERRO','erro':f'{type(e).__name__}: {e}'}
                guardar(info)
                print('FALHA ao fechar a guia Gerencial:',e)
                print('Controle de Eventos nao sera executado sem a guia principal confirmada.')
                info['controle']={'status':'NAO_EXECUTADO','motivo':'Falha no fechamento/retorno ao SIGA'}
                info['fim']=datetime.now().isoformat(timespec='seconds')
                guardar(info)
                return 1
            guardar(info)

            # Controle de Eventos: endpoint observado no HAR devolve agregado,
            # nao as 22 colunas da planilha; por isso usar o exportador validado da v2.
            progresso(76,'Carregando Controle de Eventos')
            detalhe('baixando o Excel pela tela inicial do SIGA (download)')
            print('\n--- Controle de Eventos (download pela v2) ---')
            try:
                origem=relatorios.baixar_controle_eventos(ctx,page)
                SAIDA.mkdir(parents=True,exist_ok=True)
                destino=SAIDA/origem.name
                import shutil
                shutil.copy2(origem,destino)
                origem.unlink(missing_ok=True)  # evita duplicata na entrada antiga da v2
                info['controle']={'status':'OK','arquivo':str(destino),'metodo':'download navegador v2'}
                print('OK',destino)
            except Exception as e:
                info['controle']={'status':'ERRO','erro':f'{type(e).__name__}: {e}'}
                print('FALHA Controle:',e)
            info['fim']=datetime.now().isoformat(timespec='seconds');guardar(info)
        finally:
            ctx.close()
    erros=[x for x in info['anos'] if x['status']!='OK']
    if erros or info['controle']['status']!='OK':
        print('\nCONCLUIDO COM FALHAS. Veja:',LOG)
        return 1
    progresso(80,'Downloads do SIGA concluídos')
    print('\n7 ARQUIVOS GERADOS EM',SAIDA)
    print('IMPORTANTE: compare os arquivos com o Grupo Support antes de integrar. A v2 nao foi alterada.')
    return 0

if __name__=='__main__':
    try:sys.exit(main())
    except Exception:
        print(traceback.format_exc());sys.exit(1)
